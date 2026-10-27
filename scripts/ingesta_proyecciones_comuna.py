"""
Ingesta: Proyecciones de población por comuna 2019-2030 (Santiago de Cali).
Fuente: documento academico/Pirmides por comunas Santiago de Cali_V1.xlsx (hoja 'Proy 2019-2030')

Crea:
  - CSV microdata + ConjuntoDatos + VariableConjuntoDatos
  - DominioDesagregacion 'Grupo de edad' (tematica) con 17 desags 'X a Y'/'80 y más'
  - nivel_geografico=4 en dominio Comuna (270) y capa_geo en sus desagregaciones
  - 3 Indicadores en subtema Población (id 1) + Datos + Gráficos + GraficoSerie
  - Vinculación al ProductoAcademico id 2 (indicadores_derivados + sección 9)
"""
import csv
import os
import openpyxl
from datetime import date
from django.db import transaction
from django.core.files import File
from django.conf import settings
from indicadores.models import (
    DominioDesagregacion, Desagregacion, NivelGeografico, CapaGeo,
    Indicador, Dato, ConjuntoDatos, VariableConjuntoDatos,
    Grafico, GraficoSerie, MenuSubtema, ProductoAcademico, SeccionPaper,
)

XLSX = '/home/amiranda/Descargas/SIIED/documento academico/Pirmides por comunas Santiago de Cali_V1.xlsx'
CSV_NAME = 'microdata/Proyecciones_poblacion_comuna_2019_2030.csv'
CSV_PATH = os.path.join(settings.MEDIA_ROOT, CSV_NAME)

# ══ 1. Parsear xlsx → filas tidy ══
wb = openpyxl.load_workbook(XLSX, read_only=True)
ws = wb['Proy 2019-2030']
hdr = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
comuna_cols = [(i, h) for i, h in enumerate(hdr) if h and str(h).startswith('Comuna')]
filas = []  # (anio, comuna, sexo, edad, valor)
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[0]:
        continue
    edad, sexo, anio = row[0], str(row[1]).strip(), int(row[2])
    for i, com in comuna_cols:
        v = row[i]
        if v is not None:
            filas.append((anio, com, sexo, edad, int(v)))
print('filas tidy:', len(filas))

os.makedirs(os.path.dirname(CSV_PATH), exist_ok=True)
with open(CSV_PATH, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(['anio', 'comuna', 'sexo', 'grupo_edad', 'poblacion'])
    w.writerows(filas)
print('csv →', CSV_PATH)

with transaction.atomic():
    # ══ 2. Dominio Comuna (270): nivel + capas ══
    dom_comuna = DominioDesagregacion.objects.get(pk=270)
    dom_comuna.nivel_geografico_id = 4
    dom_comuna.save(update_fields=['nivel_geografico'])
    capas = {c.nombre.strip().lower(): c for c in CapaGeo.objects.filter(nivel_geografico_id=4)}
    linked = 0
    for d in dom_comuna.desagregaciones.all():
        capa = capas.get(d.nombre.strip().lower())
        if capa and d.capa_geo_id != capa.id:
            d.capa_geo = capa
            d.save(update_fields=['capa_geo'])
            linked += 1
    print('comuna: nivel=4, capas enlazadas:', linked)

    # ══ 3. Dominio Grupo de edad (nuevo) ══
    EDADES = ['0 a 4', '5 a 9', '10 a 14', '15 a 19', '20 a 24', '25 a 29', '30 a 34',
              '35 a 39', '40 a 44', '45 a 49', '50 a 54', '55 a 59', '60 a 64',
              '65 a 69', '70 a 74', '75 a 79', '80 y más']
    dom_edad = DominioDesagregacion.objects.create(
        nombre='Grupo de edad', tipo='tematica', en_catalogo=True)
    desag_edad = {}
    for i, e in enumerate(EDADES):
        desag_edad[e] = Desagregacion.objects.create(
            nombre=e, codigo=f'GE_{i+1:02d}', dominio=dom_edad)
    print('grupo edad dominio:', dom_edad.id, len(desag_edad), 'desags')

    dom_sexo = DominioDesagregacion.objects.get(pk=2)
    desag_sexo = {d.nombre: d for d in dom_sexo.desagregaciones.all()}
    desag_comuna = {d.nombre.upper(): d for d in dom_comuna.desagregaciones.all()}

    # ══ 4. ConjuntoDatos + variables ══
    conj = ConjuntoDatos.objects.create(
        nombre='Proyecciones de población por comuna, sexo y grupo de edad 2019-2030 — Santiago de Cali',
        descripcion='Proyecciones experimentales de población del Distrito de Santiago de Cali '
                    'desagregadas por comuna, sexo y grupo quinquenal de edad, periodo 2019-2030. '
                    'Producto del ejercicio metodológico documentado en el paper académico de proyecciones.',
        unidad_analisis='Personas', cobertura_geografica='Distrito de Santiago de Cali — comunas 1-22',
        cobertura_temporal='2019-2030', tipo_estudio='Proyección poblacional (modelo demográfico)',
        palabras_clave='población, proyección, comunas, pirámide poblacional, Santiago de Cali')
    with open(CSV_PATH, 'rb') as f:
        conj.archivo_csv.save(os.path.basename(CSV_PATH), File(f), save=True)
    VARS = [
        ('V1', 'Año de proyección', 'Año calendario al que corresponde la población proyectada (2019-2030).'),
        ('V2', 'Comuna', 'Comuna del Distrito de Santiago de Cali (1 a 22).'),
        ('V3', 'Sexo', 'Sexo de la población proyectada: Hombres o Mujeres.'),
        ('V4', 'Grupo de edad', 'Grupo quinquenal de edad: 0 a 4 años hasta 80 y más.'),
        ('V5', 'Población proyectada', 'Número de personas proyectadas para la combinación año-comuna-sexo-edad.'),
    ]
    for i, (cod, nom, defin) in enumerate(VARS):
        VariableConjuntoDatos.objects.create(
            conjunto_datos=conj, codigo_variable=cod, nombre_variable=nom,
            definicion_conceptual=defin, orden=i)
    print('conjunto:', conj.id, conj.nombre[:60])

    # ══ 5. Indicadores ══
    subtema = MenuSubtema.objects.get(pk=1)
    FUENTE = 'DANE / Proyecciones experimentales de población (metodología propia 2019-2030)'

    def mk_ind(nombre, desc, orden):
        ind = Indicador.objects.create(
            nombre=nombre, subtema=subtema, unidad_medida='Personas',
            fuente_datos=FUENTE, descripcion_indicador=desc,
            conjunto_datos=conj, nombre_conjunto_datos=conj.nombre,
            cobertura_geografica='Distrito de Santiago de Cali — comunas 1 a 22',
            periodo_referencia='2019-2030', frecuencia_difusion='Anual',
            serie_historica_disponible='2019-2030', orden=orden)
        ind.dominios_desagregacion_geografica.add(dom_comuna)
        return ind

    ind_a = mk_ind('Población por comuna de Santiago de Cali',
                   'Población total proyectada por comuna del Distrito de Santiago de Cali, 2019-2030.', 11)
    ind_b = mk_ind('Población por comuna y sexo de Santiago de Cali',
                   'Población proyectada por comuna y sexo del Distrito de Santiago de Cali, 2019-2030.', 12)
    ind_b.dominios_desagregacion_tematica.add(dom_sexo)
    ind_c = mk_ind('Población por comuna, sexo y grupo de edad de Santiago de Cali',
                   'Población proyectada por comuna, sexo y grupo quinquenal de edad del '
                   'Distrito de Santiago de Cali, 2019-2030.', 13)
    ind_c.dominios_desagregacion_tematica.add(dom_sexo, dom_edad)
    print('indicadores:', ind_a.id, ind_b.id, ind_c.id)

    # ══ 6. Datos ══
    # ind_c: un Dato por (anio, comuna, sexo, edad) → geo=comuna, tem=sexo+edad
    # ind_b: totales por (anio, comuna, sexo)        → geo=comuna, tem=sexo
    # ind_a: totales por (anio, comuna)              → geo=comuna
    from collections import defaultdict
    tot_comuna = defaultdict(int)   # (anio, desag_comuna) → total
    tot_sexo = defaultdict(int)     # (anio, desag_comuna, desag_sexo) → total
    spec = []  # (ind, anio, valor, geo_desag, [tem_desags])

    for anio, com, sexo, edad, v in filas:
        dcom = desag_comuna[com.upper()]
        dsex = desag_sexo[sexo]
        spec.append((ind_c, anio, v, dcom, [dsex, desag_edad[edad]]))
        tot_comuna[(anio, dcom)] += v
        tot_sexo[(anio, dcom, dsex)] += v
    for (anio, dcom), v in tot_comuna.items():
        spec.append((ind_a, anio, v, dcom, []))
    for (anio, dcom, dsex), v in tot_sexo.items():
        spec.append((ind_b, anio, v, dcom, [dsex]))

    objs = [Dato(indicador=ind, fecha_dato=date(anio, 1, 1), valor_dato=v)
            for ind, anio, v, geo, tem in spec]
    Dato.objects.bulk_create(objs, batch_size=1000)
    print('datos creados:', len(objs), '| esperados:',
          8976 + len(tot_comuna) + len(tot_sexo))

    Through = Dato.desagregaciones_geograficas.through
    ThroughT = Dato.desagregaciones_tematicas.through
    geo_rows, tem_rows = [], []
    for dato, (ind, anio, v, geo, tem) in zip(objs, spec):
        if geo is not None:
            geo_rows.append(Through(dato_id=dato.id, desagregacion_id=geo.id))
        for t in tem:
            tem_rows.append(ThroughT(dato_id=dato.id, desagregacion_id=t.id))
    Through.objects.bulk_create(geo_rows, batch_size=1000)
    ThroughT.objects.bulk_create(tem_rows, batch_size=1000)
    print('m2m geo:', len(geo_rows), '| m2m tem:', len(tem_rows))

    # ══ 7. Gráficos + series ══
    g_a = Grafico.objects.create(
        indicador=ind_a, nombre=ind_a.nombre, tipo_grafico='bar',
        titulo='Población proyectada por comuna', eje_x_titulo='Año',
        eje_y_titulo='Personas', posicion_leyenda='right')
    for i, com in enumerate(sorted(desag_comuna, key=lambda n: int(n.split()[-1]))):
        s = GraficoSerie.objects.create(grafico=g_a, orden=i,
                                        nombre_personalizado=com.title())
        s.desagregaciones.add(desag_comuna[com])
        s.sincronizar_fks_legacy()

    g_b = Grafico.objects.create(
        indicador=ind_b, nombre=ind_b.nombre, tipo_grafico='line',
        titulo='Población proyectada por comuna y sexo', eje_x_titulo='Año',
        eje_y_titulo='Personas', posicion_leyenda='right')
    for i, sx in enumerate(['Hombres', 'Mujeres']):
        s = GraficoSerie.objects.create(grafico=g_b, orden=i,
                                        nombre_personalizado=sx)
        s.desagregaciones.add(desag_sexo[sx])
        s.sincronizar_fks_legacy()

    Grafico.objects.create(
        indicador=ind_c, nombre=ind_c.nombre,
        tipo_grafico='piramide_poblacional', titulo='Pirámide poblacional por comuna',
        eje_y_titulo='Grupo de edad', posicion_leyenda='right')
    print('gráficos:', g_a.id, g_b.id, '(pirámide sin series, patrón ind 3)')

    # ══ 8. Vincular al ProductoAcademico 2 ══
    paper = ProductoAcademico.objects.get(pk=2)
    paper.indicadores_derivados.add(ind_a, ind_b, ind_c)
    sec9 = SeccionPaper.objects.filter(producto=paper, numero_orden=9).first()
    if sec9:
        sec9.indicador_vinculado = ind_c
        sec9.save(update_fields=['indicador_vinculado'])
    print('paper 2: indicadores_derivados +=', [ind_a.id, ind_b.id, ind_c.id])

print('══ INGESTA COMPLETA ══')

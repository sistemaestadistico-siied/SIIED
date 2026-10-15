from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.db.models import Avg, Count, Q
from django.urls import reverse
from django.utils import timezone
from datetime import timedelta
import json
import csv

from .models import (
    Organismo,
    Suborganismo,
    EquipoOrganismo,
    RespuestaSatisfaccionRapida, 
    RespuestaSatisfaccionCompleta
)
from .forms_satisfaccion import (
    RespuestaSatisfaccionCompletaForm,
    OrganismoForm
)


def get_client_ip(request):
    """Obtiene la IP del cliente"""
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        ip = x_forwarded_for.split(',')[0]
    else:
        ip = request.META.get('REMOTE_ADDR')
    return ip


def satisfaccion_rapida(request, codigo_organismo, codigo_equipo=None):
    """
    LEGACY: enlaces de banners antiguos (/rapida/?r=<1-5>) que siguen
    circulando en correos ya enviados.

    La estrategia actual es un banner honesto con CTA único al formulario
    completo, así que esta vista solo registra la calificación del clic
    (dato real) y redirige al formulario. Se conserva ?ampliar= para
    vincular la respuesta completa con la rápida si el usuario la envía.
    """
    organismo = get_object_or_404(Organismo, codigo=codigo_organismo, activo=True)

    # Obtener equipo si se proporciona código
    equipo = None
    if codigo_equipo:
        equipo = get_object_or_404(
            EquipoOrganismo,
            suborganismo__organismo=organismo,
            codigo=codigo_equipo,
            activo=True
        )

    # Registrar la calificación del clic si es válida
    calificacion = request.GET.get('r')
    respuesta = None
    if calificacion and calificacion.isdigit() and int(calificacion) in range(1, 6):
        respuesta = RespuestaSatisfaccionRapida.objects.create(
            organismo=organismo,
            equipo=equipo,
            calificacion=int(calificacion),
            ip_address=get_client_ip(request),
            user_agent=request.META.get('HTTP_USER_AGENT', '')[:500]
        )

    # Redirigir al formulario completo
    if equipo:
        url = reverse('indicadores:satisfaccion_formulario_equipo', args=[organismo.codigo, equipo.codigo])
    else:
        url = reverse('indicadores:satisfaccion_formulario', args=[organismo.codigo])

    if respuesta:
        url += f'?ampliar={respuesta.id}'

    return redirect(url)


def satisfaccion_formulario(request, codigo_organismo, codigo_equipo=None):
    """
    Vista para formulario completo de satisfacción
    URL: /satisfaccion/<codigo_organismo>/<codigo_equipo>/
    URL (legacy): /satisfaccion/<codigo_organismo>/
    """
    organismo = get_object_or_404(Organismo, codigo=codigo_organismo, activo=True)
    
    # Obtener equipo si se proporciona código
    equipo = None
    if codigo_equipo:
        equipo = get_object_or_404(
            EquipoOrganismo, 
            suborganismo__organismo=organismo, 
            codigo=codigo_equipo, 
            activo=True
        )
    
    # Verificar si viene de una respuesta rápida (desde URL o POST)
    respuesta_rapida_id = request.GET.get('ampliar') or request.POST.get('respuesta_rapida_id')
    respuesta_rapida = None
    
    if respuesta_rapida_id:
        try:
            respuesta_rapida = RespuestaSatisfaccionRapida.objects.get(
                id=respuesta_rapida_id,
                organismo=organismo,
                ampliada=False
            )
            # Si la respuesta rápida tiene equipo, usarlo
            if respuesta_rapida.equipo:
                equipo = respuesta_rapida.equipo
        except RespuestaSatisfaccionRapida.DoesNotExist:
            pass
    
    if request.method == 'POST':
        form = RespuestaSatisfaccionCompletaForm(request.POST)

        if form.is_valid():
            respuesta_completa = form.save(commit=False)
            respuesta_completa.organismo = organismo
            respuesta_completa.equipo = equipo
            respuesta_completa.ip_address = get_client_ip(request)
            respuesta_completa.user_agent = request.META.get('HTTP_USER_AGENT', '')[:500]
            
            # Vincular con respuesta rápida si existe
            if respuesta_rapida:
                respuesta_completa.respuesta_rapida = respuesta_rapida
                respuesta_rapida.ampliada = True
                respuesta_rapida.save()
            
            respuesta_completa.save()
            
            # Redirigir a página de agradecimiento con parámetros
            if equipo:
                url_name = 'indicadores:satisfaccion_gracias_completa_equipo'
                return redirect(f"{reverse(url_name, args=[organismo.codigo, equipo.codigo])}?promedio={respuesta_completa.calcular_promedio_satisfaccion()}")
            else:
                url_name = 'indicadores:satisfaccion_gracias_completa'
                return redirect(f"{reverse(url_name, args=[organismo.codigo])}?promedio={respuesta_completa.calcular_promedio_satisfaccion()}")
    else:
        # Formulario sin pre-llenar para evitar sesgo estadístico
        form = RespuestaSatisfaccionCompletaForm()
    
    context = {
        'form': form,
        'organismo': organismo,
        'equipo': equipo,
        'respuesta_rapida': respuesta_rapida,
    }
    
    return render(request, 'satisfaccion/formulario.html', context)


def satisfaccion_gracias_completa(request, codigo_organismo, codigo_equipo=None):
    """
    Vista para página de agradecimiento después de completar el formulario
    URL: /satisfaccion/<codigo_organismo>/<codigo_equipo>/gracias/
    URL (legacy): /satisfaccion/<codigo_organismo>/gracias/
    """
    organismo = get_object_or_404(Organismo, codigo=codigo_organismo, activo=True)
    promedio = request.GET.get('promedio', 0)
    
    context = {
        'organismo': organismo,
        'promedio': promedio,
    }
    
    return render(request, 'satisfaccion/gracias_completa.html', context)


@staff_member_required
def satisfaccion_dashboard(request):
    """
    Dashboard de estadísticas de satisfacción
    URL: /admin/satisfaccion/dashboard/
    """
    # Filtros
    organismo_id = request.GET.get('organismo')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')
    
    # Query base
    respuestas_rapidas = RespuestaSatisfaccionRapida.objects.all()
    respuestas_completas = RespuestaSatisfaccionCompleta.objects.all()
    
    # Aplicar filtros
    if organismo_id:
        respuestas_rapidas = respuestas_rapidas.filter(organismo_id=organismo_id)
        respuestas_completas = respuestas_completas.filter(organismo_id=organismo_id)
    
    if fecha_desde:
        respuestas_rapidas = respuestas_rapidas.filter(fecha_registro__gte=fecha_desde)
        respuestas_completas = respuestas_completas.filter(fecha_registro__gte=fecha_desde)
    
    if fecha_hasta:
        respuestas_rapidas = respuestas_rapidas.filter(fecha_registro__lte=fecha_hasta)
        respuestas_completas = respuestas_completas.filter(fecha_registro__lte=fecha_hasta)
    
    # Métricas principales - SEPARADAS (no son comparables)
    total_respuestas_rapidas = respuestas_rapidas.count()
    total_respuestas_completas = respuestas_completas.count()
    
    # Promedio de respuestas rápidas (escala 1-5)
    promedio_rapidas = respuestas_rapidas.aggregate(Avg('calificacion'))['calificacion__avg'] or 0
    
    # Promedio de respuestas completas (escala 1-5)
    promedios_completas = []
    for respuesta in respuestas_completas:
        promedios_completas.append(respuesta.calcular_promedio_satisfaccion())
    promedio_completas = sum(promedios_completas) / len(promedios_completas) if promedios_completas else 0
    
    # Respuestas últimos 30 días - SEPARADAS
    hace_30_dias = timezone.now() - timedelta(days=30)
    respuestas_rapidas_30_dias = respuestas_rapidas.filter(fecha_registro__gte=hace_30_dias).count()
    respuestas_completas_30_dias = respuestas_completas.filter(fecha_registro__gte=hace_30_dias).count()
    
    # Distribución de calificaciones (respuestas rápidas)
    distribucion_rapidas = respuestas_rapidas.values('calificacion').annotate(
        count=Count('id')
    ).order_by('calificacion')
    
    # Top organismos por número de respuestas RÁPIDAS
    top_organismos_rapidas = Organismo.objects.annotate(
        total_rapidas=Count('respuestas_rapidas_organismo')
    ).filter(total_rapidas__gt=0).order_by('-total_rapidas')[:10]
    
    # Top organismos por número de respuestas COMPLETAS
    top_organismos_completas = Organismo.objects.annotate(
        total_completas=Count('respuestas_completas_organismo')
    ).filter(total_completas__gt=0).order_by('-total_completas')[:10]
    
    # Métricas por equipos - RESPUESTAS RÁPIDAS
    equipos_rapidas = EquipoOrganismo.objects.annotate(
        total_rapidas=Count('respuestas_rapidas'),
        promedio_calificacion=Avg('respuestas_rapidas__calificacion')
    ).filter(total_rapidas__gt=0).select_related('suborganismo__organismo').order_by('-total_rapidas')[:15]
    
    # Métricas por equipos - RESPUESTAS COMPLETAS
    equipos_completas = EquipoOrganismo.objects.annotate(
        total_completas=Count('respuestas_completas')
    ).filter(total_completas__gt=0).select_related('suborganismo__organismo').order_by('-total_completas')[:15]
    
    # Datos para gráficos (JSON)
    # Gráfico 1: Distribución de calificaciones
    distribucion_data = {
        'labels': ['Muy insatisfecho', 'Insatisfecho', 'Neutral', 'Satisfecho', 'Muy satisfecho'],
        'data': [0, 0, 0, 0, 0]
    }
    for item in distribucion_rapidas:
        distribucion_data['data'][item['calificacion'] - 1] = item['count']
    
    # Gráfico 2: Promedios por dimensión (solo respuestas completas)
    if respuestas_completas.exists():
        promedio_utilidad = respuestas_completas.aggregate(
            Avg('utilidad_tiempo_respuesta'),
            Avg('utilidad_claridad')
        )
        promedio_relevancia = respuestas_completas.aggregate(
            Avg('relevancia_periodicidad'),
            Avg('relevancia_concordancia')
        )
        promedio_accesibilidad = respuestas_completas.aggregate(
            Avg('accesibilidad_canal'),
            Avg('accesibilidad_formato')
        )
        
        dimensiones_data = {
            'labels': ['Utilidad', 'Relevancia', 'Accesibilidad'],
            'data': [
                round((
                    (promedio_utilidad['utilidad_tiempo_respuesta__avg'] or 0) +
                    (promedio_utilidad['utilidad_claridad__avg'] or 0)
                ) / 2, 2),
                round((
                    (promedio_relevancia['relevancia_periodicidad__avg'] or 0) +
                    (promedio_relevancia['relevancia_concordancia__avg'] or 0)
                ) / 2, 2),
                round((
                    (promedio_accesibilidad['accesibilidad_canal__avg'] or 0) +
                    (promedio_accesibilidad['accesibilidad_formato__avg'] or 0)
                ) / 2, 2),
            ]
        }
    else:
        dimensiones_data = {
            'labels': ['Utilidad', 'Relevancia', 'Accesibilidad'],
            'data': [0, 0, 0]
        }
    
    
    # ========================================================================
    # INDICADORES ESPECÍFICOS PARA RESPUESTAS COMPLETAS
    # ========================================================================
    
    # Variables de las 3 dimensiones
    variables = {
        'utilidad_tiempo_respuesta': 'Utilidad: Tiempo de respuesta',
        'utilidad_claridad': 'Utilidad: Claridad, precisión y facilidad de uso',
        'relevancia_periodicidad': 'Relevancia: Correspondencia de periodicidad',
        'relevancia_concordancia': 'Relevancia: Concordancia entre lo solicitado y entregado',
        'accesibilidad_canal': 'Accesibilidad: Canal de solicitud',
        'accesibilidad_formato': 'Accesibilidad: Formato de entrega',
    }
    
    # INDICADOR 1: Porcentaje de satisfacción por variable y categoría (SV)
    indicador_sv = {}
    for campo, nombre in variables.items():
        total_var = respuestas_completas.filter(**{f'{campo}__isnull': False}).count()
        if total_var > 0:
            distribucion = {}
            for categoria in range(1, 6):
                count = respuestas_completas.filter(**{campo: categoria}).count()
                porcentaje = (count / total_var) * 100
                distribucion[categoria] = {
                    'count': count,
                    'porcentaje': round(porcentaje, 2)
                }
            indicador_sv[campo] = {
                'nombre': nombre,
                'total': total_var,
                'distribucion': distribucion
            }
    
    # INDICADOR 2: Índice Neto de Satisfacción por Dimensión (INSD)
    indicador_insd = {}
    
    dimensiones_campos = {
        'Utilidad': ['utilidad_tiempo_respuesta', 'utilidad_claridad'],
        'Relevancia': ['relevancia_periodicidad', 'relevancia_concordancia'],
        'Accesibilidad': ['accesibilidad_canal', 'accesibilidad_formato']
    }
    
    for dimension, campos in dimensiones_campos.items():
        # Contar todas las respuestas de esta dimensión (ambas variables)
        total_dimension = 0
        satisfechas = 0
        no_satisfechas = 0
        
        for campo in campos:
            # Satisfechas: categorías 4 y 5
            satisfechas += respuestas_completas.filter(**{f'{campo}__in': [4, 5]}).count()
            # No satisfechas: categorías 1, 2 y 3
            no_satisfechas += respuestas_completas.filter(**{f'{campo}__in': [1, 2, 3]}).count()
            # Total
            total_dimension += respuestas_completas.filter(**{f'{campo}__isnull': False}).count()
        
        if total_dimension > 0:
            diferencia = satisfechas - no_satisfechas
            insd_valor = abs(diferencia) / total_dimension * 100
            percepcion = 'positiva' if diferencia > 0 else 'negativa' if diferencia < 0 else 'neutral'
            
            indicador_insd[dimension] = {
                'satisfechas': satisfechas,
                'no_satisfechas': no_satisfechas,
                'total': total_dimension,
                'diferencia': diferencia,
                'insd': round(insd_valor, 2),
                'percepcion': percepcion
            }
    
    # INDICADOR 3: Índice Neto de Satisfacción Global (INSG)
    total_global = 0
    satisfechas_global = 0
    no_satisfechas_global = 0
    
    for campos in dimensiones_campos.values():
        for campo in campos:
            satisfechas_global += respuestas_completas.filter(**{f'{campo}__in': [4, 5]}).count()
            no_satisfechas_global += respuestas_completas.filter(**{f'{campo}__in': [1, 2, 3]}).count()
            total_global += respuestas_completas.filter(**{f'{campo}__isnull': False}).count()
    
    if total_global > 0:
        diferencia_global = satisfechas_global - no_satisfechas_global
        insg_valor = abs(diferencia_global) / total_global * 100
        percepcion_global = 'positiva' if diferencia_global > 0 else 'negativa' if diferencia_global < 0 else 'neutral'
        
        indicador_insg = {
            'satisfechas': satisfechas_global,
            'no_satisfechas': no_satisfechas_global,
            'total': total_global,
            'diferencia': diferencia_global,
            'insg': round(insg_valor, 2),
            'percepcion': percepcion_global
        }
    else:
        indicador_insg = None
    
    # Últimas respuestas para tablas
    ultimas_rapidas = respuestas_rapidas.select_related('organismo').order_by('-fecha_registro')[:10]
    ultimas_completas = respuestas_completas.select_related('organismo').order_by('-fecha_registro')[:10]
    
    context = {
        'total_respuestas_rapidas': total_respuestas_rapidas,
        'total_respuestas_completas': total_respuestas_completas,
        'promedio_rapidas': round(promedio_rapidas, 2),
        'promedio_completas': round(promedio_completas, 2),
        'respuestas_rapidas_30_dias': respuestas_rapidas_30_dias,
        'respuestas_completas_30_dias': respuestas_completas_30_dias,
        'top_organismos_rapidas': top_organismos_rapidas,
        'top_organismos_completas': top_organismos_completas,
        'equipos_rapidas': equipos_rapidas,
        'equipos_completas': equipos_completas,
        'organismos': Organismo.objects.filter(activo=True).order_by('nombre'),
        'organismo_seleccionado': organismo_id,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'distribucion_data': json.dumps(distribucion_data),
        'dimensiones_data': json.dumps(dimensiones_data),
        'ultimas_rapidas': ultimas_rapidas,
        'ultimas_completas': ultimas_completas,
        # Indicadores específicos para respuestas completas
        'indicador_sv': indicador_sv,
        'indicador_insd': indicador_insd,
        'indicador_insg': indicador_insg,
    }
    
    return render(request, 'admin/satisfaccion_dashboard.html', context)


@staff_member_required
def satisfaccion_organismos(request):
    """
    Vista para gestionar organismos y generar URLs
    URL: /admin/satisfaccion/organismos/
    """
    if request.method == 'POST':
        form = OrganismoForm(request.POST)
        if form.is_valid():
            organismo = form.save()
            messages.success(request, f'Organismo "{organismo.nombre}" creado exitosamente.')
            return redirect('satisfaccion_organismos')
    else:
        form = OrganismoForm()
    
    organismos = Organismo.objects.all().order_by('-fecha_creacion')
    
    # Generar URLs completas para cada organismo
    for organismo in organismos:
        organismo.url_rapida = request.build_absolute_uri(
            organismo.get_url_satisfaccion_rapida()
        )
        organismo.url_completa = request.build_absolute_uri(
            organismo.get_url_formulario_completo()
        )
    
    context = {
        'form': form,
        'organismos': organismos,
    }
    
    return render(request, 'admin/satisfaccion_organismos.html', context)


@staff_member_required
def satisfaccion_organismo_editar(request, organismo_id):
    """
    Editar un organismo existente
    """
    organismo = get_object_or_404(Organismo, id=organismo_id)
    
    if request.method == 'POST':
        form = OrganismoForm(request.POST, instance=organismo)
        if form.is_valid():
            form.save()
            messages.success(request, f'Organismo "{organismo.nombre}" actualizado exitosamente.')
            return redirect('satisfaccion_organismos')
    else:
        form = OrganismoForm(instance=organismo)
    
    context = {
        'form': form,
        'organismo': organismo,
        'editando': True,
    }
    
    return render(request, 'admin/satisfaccion_organismo_form.html', context)


def _generar_banner_v3(organismo, etiqueta, url_formulario):
    """
    Banner de satisfacción (diseño V3): un único CTA que lleva
    directamente al formulario completo.

    Diseño tabular con estilos en línea, optimizado para copiar/pegarse
    en Gmail, Google Docs y Word sin bordes ni truncamiento.
    """
    # Estilos anti-borde para Word y compatibilidad con clientes de email
    tbl = 'border="0" cellpadding="0" cellspacing="0" style="border-collapse: collapse; border: none; mso-table-lspace:0pt; mso-table-rspace:0pt;"'
    td_base = 'border: none; padding: 0; vertical-align: middle;'

    return f'''<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Califique nuestra atención</title>
</head>
<body style="margin: 0; padding: 0; font-family: Arial, sans-serif; background-color: #ffffff;">
    <!-- Wrapper fluido 100% para evitar truncamiento en Docs/Word -->
    <table {tbl} width="100%">
        <tr>
            <td align="left" style="{td_base}">
                <!-- Banner V3: ancho fijo 520px — igual a la firma de Gmail -->
                <table {tbl} width="520" style="max-width: 520px;">
                    <tr>
                        <td style="background-color: #10184a; border-radius: 16px; padding: 24px 30px 22px;">
                            <!-- Fila encabezado: organismo + etiqueta -->
                            <table {tbl} width="100%">
                                <tr>
                                    <td style="{td_base} color: #c9cfe8; font-size: 12px; font-weight: bold; letter-spacing: 0.2px;">
                                        {organismo.nombre}
                                    </td>
                                    <td align="right" style="{td_base}">
                                        <span style="background-color: #2f3d9c; color: #ffffff; padding: 4px 10px; border-radius: 999px; font-size: 10px; font-weight: bold; letter-spacing: 0.6px; white-space: nowrap;">
                                            {etiqueta}
                                        </span>
                                    </td>
                                </tr>
                            </table>

                            <!-- Titular -->
                            <p style="color: #ffffff; margin: 20px 0 12px 0; font-size: 21px; font-weight: bold; line-height: 1.3;">
                                Tu opinión nos ayuda a <span style="color: #ffb648;">mejorar</span>
                            </p>

                            <!-- Subtexto -->
                            <p style="color: #dfe3f5; margin: 0 0 20px 0; font-size: 13px; line-height: 1.5;">
                                Ahora queremos conocer tu satisfacción con la información estadística recibida.
                            </p>

                            <!-- CTA único -->
                            <table {tbl}>
                                <tr>
                                    <td style="background-color: #ffb648; border-radius: 999px;">
                                        <a href="{url_formulario}" style="display: inline-block; padding: 11px 22px; color: #1a1440; font-size: 14px; font-weight: bold; text-decoration: none;">
                                            Califica tu experiencia →
                                        </a>
                                    </td>
                                </tr>
                            </table>

                            <!-- Pie -->
                            <p style="color: #8f97c4; margin: 12px 0 0 0; font-size: 10.5px;">
                                Toma menos de 2 minutos
                            </p>
                        </td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</body>
</html>'''


def generar_codigo_email_equipo(organismo, equipo, request):
    """
    Genera el código HTML del banner de satisfacción para un equipo específico.
    El banner enlaza directamente al formulario completo del equipo.
    """
    url_formulario = request.build_absolute_uri(
        reverse('indicadores:satisfaccion_formulario_equipo', args=[organismo.codigo, equipo.codigo])
    )
    return _generar_banner_v3(organismo, equipo.nombre, url_formulario)


def satisfaccion_instructivo(request):
    """
    Instructivo interactivo para empleados sobre cómo embeber el código en emails
    Flujo: Organismo → Suborganismo (obligatorio) → Equipo
    URL: /satisfaccion/instructivo/
    """
    # Cargar organismos activos
    organismos = Organismo.objects.filter(activo=True).order_by('nombre')
    
    # Obtener selecciones del usuario
    organismo_id = request.GET.get('organismo')
    suborganismo_id = request.GET.get('suborganismo')
    equipo_id = request.GET.get('equipo')
    
    organismo_seleccionado = None
    suborganismo_seleccionado = None
    equipo_seleccionado = None
    suborganismos_disponibles = []
    equipos_disponibles = []
    codigo_html = None
    
    # Paso 1: Organismo seleccionado
    if organismo_id:
        try:
            organismo_seleccionado = Organismo.objects.get(id=organismo_id, activo=True)
            # Cargar suborganismos del organismo
            suborganismos_disponibles = Suborganismo.objects.filter(
                organismo=organismo_seleccionado,
                activo=True
            ).order_by('orden', 'nombre')
            
            # Paso 2: Suborganismo seleccionado
            if suborganismo_id:
                try:
                    suborganismo_seleccionado = suborganismos_disponibles.get(id=suborganismo_id)
                    # Cargar equipos del suborganismo
                    equipos_disponibles = EquipoOrganismo.objects.filter(
                        suborganismo=suborganismo_seleccionado,
                        activo=True
                    ).order_by('nombre')
                    
                    # Paso 3: Equipo seleccionado
                    if equipo_id:
                        try:
                            equipo_seleccionado = equipos_disponibles.get(id=equipo_id)
                            # Generar el código HTML del email para el equipo
                            codigo_html = generar_codigo_email_equipo(organismo_seleccionado, equipo_seleccionado, request)
                        except EquipoOrganismo.DoesNotExist:
                            pass
                            
                except Suborganismo.DoesNotExist:
                    pass
                    
        except Organismo.DoesNotExist:
            pass
    
    context = {
        'organismos': organismos,
        'organismo_seleccionado': organismo_seleccionado,
        'suborganismos_disponibles': suborganismos_disponibles,
        'suborganismo_seleccionado': suborganismo_seleccionado,
        'equipos_disponibles': equipos_disponibles,
        'equipo_seleccionado': equipo_seleccionado,
        'codigo_html': codigo_html,
    }
    
    return render(request, 'satisfaccion/instructivo.html', context)


def generar_codigo_email(organismo, request):
    """
    Genera el código HTML del banner de satisfacción para un organismo
    (sin equipo). El banner enlaza directamente al formulario completo.
    """
    url_formulario = request.build_absolute_uri(
        reverse('indicadores:satisfaccion_formulario', args=[organismo.codigo])
    )
    return _generar_banner_v3(organismo, 'SIIED', url_formulario)


@staff_member_required
def satisfaccion_exportar(request):
    """
    Exportar datos de satisfacción a Excel
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from django.http import HttpResponse
    from datetime import datetime
    
    # Crear workbook
    wb = openpyxl.Workbook()
    
    # Hoja 1: Respuestas Rápidas
    ws1 = wb.active
    ws1.title = "Respuestas Rápidas"
    
    # Encabezados
    headers1 = ['ID', 'Organismo', 'Calificación', 'Fecha', 'IP', 'Ampliada']
    ws1.append(headers1)
    
    # Estilo de encabezados
    for cell in ws1[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
    
    # Datos
    for respuesta in RespuestaSatisfaccionRapida.objects.all().order_by('-fecha_registro'):
        ws1.append([
            respuesta.id,
            respuesta.organismo.nombre,
            respuesta.get_calificacion_display(),
            respuesta.fecha_registro.strftime('%Y-%m-%d %H:%M'),
            respuesta.ip_address or '',
            'Sí' if respuesta.ampliada else 'No'
        ])
    
    # Hoja 2: Respuestas Completas
    ws2 = wb.create_sheet("Respuestas Completas")
    
    headers2 = [
        'ID', 'Organismo', 'Fecha Respuesta', 'Fecha Registro',
        'Utilidad Tiempo', 'Utilidad Claridad',
        'Relevancia Periodicidad', 'Relevancia Concordancia',
        'Accesibilidad Canal', 'Accesibilidad Formato',
        'Promedio General', 'Razón Insatisfacción'
    ]
    ws2.append(headers2)
    
    # Estilo de encabezados
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
    
    # Datos
    for respuesta in RespuestaSatisfaccionCompleta.objects.all().order_by('-fecha_registro'):
        ws2.append([
            respuesta.id,
            respuesta.organismo.nombre,
            respuesta.fecha_respuesta_solicitud.strftime('%Y-%m-%d'),
            respuesta.fecha_registro.strftime('%Y-%m-%d %H:%M'),
            respuesta.utilidad_tiempo_respuesta,
            respuesta.utilidad_claridad,
            respuesta.relevancia_periodicidad,
            respuesta.relevancia_concordancia,
            respuesta.accesibilidad_canal,
            respuesta.accesibilidad_formato,
            round(respuesta.calcular_promedio_satisfaccion(), 2),
            respuesta.razon_insatisfaccion or ''
        ])
    
    # Ajustar anchos de columna
    for ws in [ws1, ws2]:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    # Preparar respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'satisfaccion_ciudadana_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response


@staff_member_required
def exportar_respuestas_rapidas_csv(request):
    """
    Exportar respuestas rápidas a CSV con filtros del dashboard
    """
    # Aplicar los mismos filtros que el dashboard
    organismo_id = request.GET.get('organismo')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')
    
    # Convertir strings 'None' a None real
    if organismo_id == 'None' or not organismo_id:
        organismo_id = None
    if fecha_desde == 'None' or not fecha_desde:
        fecha_desde = None
    if fecha_hasta == 'None' or not fecha_hasta:
        fecha_hasta = None
    
    respuestas = RespuestaSatisfaccionRapida.objects.select_related('organismo', 'equipo').all()
    
    if organismo_id:
        respuestas = respuestas.filter(organismo_id=organismo_id)
    if fecha_desde:
        respuestas = respuestas.filter(fecha_registro__gte=fecha_desde)
    if fecha_hasta:
        respuestas = respuestas.filter(fecha_registro__lte=fecha_hasta)
    
    respuestas = respuestas.order_by('-fecha_registro')
    
    # Crear respuesta CSV
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="respuestas_rapidas_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv"'
    response.write('\ufeff')  # BOM para Excel
    
    writer = csv.writer(response)
    writer.writerow([
        'ID',
        'Organismo',
        'Equipo',
        'Calificación',
        'Calificación (texto)',
        'Fecha Respuesta Solicitud',
        'Fecha Registro',
        'Ampliada',
        'IP',
        'Comentarios'
    ])
    
    for r in respuestas:
        writer.writerow([
            r.id,
            r.organismo.nombre if r.organismo else '',
            r.equipo.nombre if r.equipo else '',
            r.calificacion,
            r.get_calificacion_display(),
            r.fecha_respuesta_solicitud.strftime('%Y-%m-%d') if r.fecha_respuesta_solicitud else '',
            r.fecha_registro.strftime('%Y-%m-%d %H:%M:%S'),
            'Sí' if r.ampliada else 'No',
            r.ip_address or '',
            r.comentarios_adicionales or ''
        ])
    
    return response


@staff_member_required
def exportar_respuestas_completas_csv(request):
    """
    Exportar respuestas completas a CSV con filtros del dashboard
    """
    # Aplicar los mismos filtros que el dashboard
    organismo_id = request.GET.get('organismo')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')
    
    # Convertir strings 'None' a None real
    if organismo_id == 'None' or not organismo_id:
        organismo_id = None
    if fecha_desde == 'None' or not fecha_desde:
        fecha_desde = None
    if fecha_hasta == 'None' or not fecha_hasta:
        fecha_hasta = None
    
    respuestas = RespuestaSatisfaccionCompleta.objects.select_related('organismo', 'equipo').all()
    
    if organismo_id:
        respuestas = respuestas.filter(organismo_id=organismo_id)
    if fecha_desde:
        respuestas = respuestas.filter(fecha_registro__gte=fecha_desde)
    if fecha_hasta:
        respuestas = respuestas.filter(fecha_registro__lte=fecha_hasta)
    
    respuestas = respuestas.order_by('-fecha_registro')
    
    # Crear respuesta CSV
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="respuestas_completas_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv"'
    response.write('\ufeff')  # BOM para Excel
    
    writer = csv.writer(response)
    writer.writerow([
        'ID',
        'Organismo',
        'Equipo',
        'Fecha Respuesta Solicitud',
        'Fecha Registro',
        'Utilidad: Tiempo Respuesta',
        'Utilidad: Claridad',
        'Relevancia: Periodicidad',
        'Relevancia: Concordancia',
        'Accesibilidad: Canal',
        'Accesibilidad: Formato',
        'Razón Insatisfacción',
        'IP'
    ])
    
    for r in respuestas:
        writer.writerow([
            r.id,
            r.organismo.nombre if r.organismo else '',
            r.equipo.nombre if r.equipo else '',
            r.fecha_respuesta_solicitud.strftime('%Y-%m-%d') if r.fecha_respuesta_solicitud else '',
            r.fecha_registro.strftime('%Y-%m-%d %H:%M:%S'),
            r.utilidad_tiempo_respuesta or '',
            r.utilidad_claridad or '',
            r.relevancia_periodicidad or '',
            r.relevancia_concordancia or '',
            r.accesibilidad_canal or '',
            r.accesibilidad_formato or '',
            r.razon_insatisfaccion or '',
            r.ip_address or ''
        ])

    return response


def satisfaccion_banners_suborganismo(request):
    """
    Vista de banners en dos pasos:
    1. Seleccionar organismo
    2. Seleccionar suborganismo → ver banners de sus equipos
    URL: /satisfaccion/banners/
    """
    organismos = Organismo.objects.filter(activo=True).order_by('nombre')

    organismo_id = request.GET.get('organismo')
    suborganismo_id = request.GET.get('suborganismo')

    organismo_seleccionado = None
    suborganismo_seleccionado = None
    suborganismos = []
    equipos_con_banner = []

    if organismo_id:
        try:
            organismo_seleccionado = Organismo.objects.get(id=organismo_id, activo=True)
            suborganismos = Suborganismo.objects.filter(
                organismo=organismo_seleccionado, activo=True
            ).order_by('orden', 'nombre')
        except Organismo.DoesNotExist:
            pass

    if suborganismo_id and organismo_seleccionado:
        try:
            suborganismo_seleccionado = Suborganismo.objects.select_related('organismo').get(
                id=suborganismo_id, activo=True, organismo=organismo_seleccionado
            )
            equipos = EquipoOrganismo.objects.filter(
                suborganismo=suborganismo_seleccionado, activo=True
            ).order_by('nombre')

            for equipo in equipos:
                equipos_con_banner.append({
                    'equipo': equipo,
                    'banner_html': generar_codigo_email_equipo(organismo_seleccionado, equipo, request),
                })
        except Suborganismo.DoesNotExist:
            pass

    context = {
        'organismos': organismos,
        'organismo_seleccionado': organismo_seleccionado,
        'suborganismos': suborganismos,
        'suborganismo_seleccionado': suborganismo_seleccionado,
        'equipos_con_banner': equipos_con_banner,
    }
    return render(request, 'satisfaccion/banners_por_suborganismo.html', context)

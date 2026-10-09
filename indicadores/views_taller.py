from .decorators import delegado_required, taller_staff_required
from django.db.models import Count
from django.http import FileResponse, Http404, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.views.decorators.http import require_GET

from . import wizard_taller as wz
from .models import Organismo, Suborganismo
from .models_taller import FuenteRequerida, InscripcionTaller as Inscripcion, RetoAnalisis
from .models_catalogo import EvaluacionSeguridad, FuenteDatos, VariableFuente

SESSION_KEY = 'inscripcion_id'
FUENTE_KEY = 'fuente_activa_id'

# Pasos de F1 que editan una FuenteDatos individual
PASOS_FUENTE = {'fuente-actualizacion', 'fuente-rol', 'fuente-proposito',
                'fuente-variables', 'fuente-calidad', 'fuente-diccionario'}
# Pasos de F3: editan la EvaluacionSeguridad de la fuente activa (F3 es por fuente)
PASOS_SEGURIDAD = {'seguridad-datos', 'seguridad-uso',
                   'seguridad-divulgacion', 'seguridad-valoracion'}
# Pasos que muestran el selector de fuente
PASOS_CON_FUENTE = PASOS_FUENTE | PASOS_SEGURIDAD

# --- Visuales del grafo del ecosistema (compartidos por GEXF y SVG) ---
COLOR_NODO = {
    'gold': '#ffd700', 'catalogo': '#007bb5', 'organismo': '#f5a86b',
    'suborganismo': '#a78bfa', 'equipo': '#4fc3f7', 'conjunto': '#6fce8f',
    'reto': '#fc8181', 'fuente_informacion': '#ef92c4',
}
COLOR_ARISTA = {
    'cruce_catalogo': '#ffd700', 'matching_variable': '#6fce8f',
    'match_exacto': '#00ff88', 'match_aproximado': '#6fce8f',
    'origen_declarado': '#f5a86b', 'origen_externo': '#ef92c4',
    'requerida': '#ef92c4', 'requerida_externa': '#c4b5fd',
    'pertenencia': '#a0aec0', 'jerarquia': '#a78bfa',
    'variables_compartidas': '#38b2ac',
}
TIPO_LABEL = {
    'gold': 'Capa Gold', 'catalogo': 'Catálogo', 'organismo': 'Organismo',
    'suborganismo': 'Suborganismo', 'equipo': 'Equipo', 'conjunto': 'Conjunto',
    'reto': 'Reto de análisis', 'fuente_informacion': 'Fuente de información',
}


def _tamano_nodo(n):
    """Radio base del nodo según su tipo y métricas (misma escala que el 3D)."""
    tipo = n.get('tipo')
    if tipo == 'gold':
        return 12
    if tipo == 'catalogo':
        return 5 + min(8, (n.get('num_campos') or 0) / 30)
    if tipo == 'organismo':
        return 5 + min(8, (n.get('num_fuentes') or 0) * 1.5)
    if tipo == 'suborganismo':
        return 3.5
    if tipo == 'equipo':
        return 2.5
    if tipo == 'conjunto':
        return 2 + min(1, (n.get('num_variables') or 0) / 6)
    if tipo == 'reto':
        return 3
    return 4


def landing_page(request):
    return render(request, 'taller/landing_page.html')


def cartel(request):
    return render(request, 'taller/cartel.html')


@delegado_required
@require_GET
def api_suborganismos(request, organismo_id):
    subs = Suborganismo.objects.filter(organismo_id=organismo_id).values('id', 'nombre').order_by('nombre')
    return JsonResponse(list(subs), safe=False)


def _construir_grafo_ecosistema(solo_relacionados=False):
    """Construye nodos y aristas del grafo del ecosistema (compartido por la API y el export GEXF)."""
    import unicodedata
    from collections import Counter
    from .catalogo_datalake import get_bases, get_campos_cruzables
    from .models import (EvaluacionSeguridad, FuenteOrigen,
                         FuenteRequerida, Organismo, VariableFuente)

    def _norm(s):
        """Normaliza: lowercase + quita acentos + strip."""
        s = s.lower().strip()
        return ''.join(c for c in unicodedata.normalize('NFD', s)
                       if unicodedata.category(c) != 'Mn')

    def _norm_var(s):
        """Normaliza nombre de variable para matching: sin acentos, sin _ ni espacios."""
        s = _norm(s)
        for prefix in ('var_', 'cod_', 'num_', 'id_', 'nom_', 'desc_', 'codigo_', 'nombre_'):
            if s.startswith(prefix):
                s = s[len(prefix):]
        return s.replace('_', '').replace(' ', '').replace('-', '')

    nodos = []
    aristas = []

    # --- Nodos del catálogo (incluye Gold) ---
    bases = get_bases()
    campos_cruzables = get_campos_cruzables()
    for b in bases:
        es_gold = b['id'] == 'fact_personas'
        nodos.append({
            'id': 'cat-' + b['id'],
            'tipo': 'gold' if es_gold else 'catalogo',
            'nombre': b['nombre'],
            'descripcion': b.get('descripcion', ''),
            'organismo': b.get('organismo', ''),
            'frecuencia': b.get('frecuencia', ''),
            'cobertura': b.get('cobertura', ''),
            'num_campos': len(b.get('campos', [])),
            'campos_cruzables': sum(1 for c in b.get('campos', []) if c.get('cruzar')),
            'campos': [c['nombre'] for c in b.get('campos', [])],
        })

    # --- Aristas entre bases del catálogo (cruzar) ---
    for campo, lista_ids in campos_cruzables.items():
        for i in range(len(lista_ids)):
            for j in range(i + 1, len(lista_ids)):
                a, b2 = lista_ids[i], lista_ids[j]
                aristas.append({
                    'origen': 'cat-' + a, 'destino': 'cat-' + b2,
                    'tipo': 'cruce_catalogo', 'peso': 1, 'campo': campo,
                })

    # --- Nodos de organismos, suborganismos, equipos y conjuntos ---

    organismo_model = Organismo
    inscripciones = (Inscripcion.objects
                     .select_related('organismo', 'suborganismo')
                     .prefetch_related('fuentes__variables', 'fuentes__fuentes_origen')
                     .all())
    fuentes = (FuenteDatos.objects.select_related('inscripcion__organismo', 'inscripcion__suborganismo',
                                                'seguridad')
               .prefetch_related('variables', 'fuentes_origen').all())

    # Conteo de fuentes por organismo en una pasada (evita N+1 en el loop)
    fuentes_por_org = Counter(f.inscripcion.organismo_id for f in fuentes)

    # Índice de campos del catálogo normalizados para matching de variables
    campos_catalogo_norm = {}  # norm_var -> [cat_ids]
    for b in bases:
        for c in b.get('campos', []):
            campos_catalogo_norm.setdefault(_norm_var(c['nombre']), []).append('cat-' + b['id'])

    # Determinar organismos y suborganismos que tienen al menos una inscripción
    org_ids_con_insc = set(inscripciones.values_list('organismo_id', flat=True))
    sub_ids_con_insc = set(inscripciones.exclude(suborganismo__isnull=True)
                           .values_list('suborganismo_id', flat=True))

    if solo_relacionados:
        organismos_qs = organismo_model.objects.filter(pk__in=org_ids_con_insc)
    else:
        organismos_qs = organismo_model.objects.all()
    organismos_qs = organismos_qs.prefetch_related('suborganismos')

    for org in organismos_qs:
        nodos.append({
            'id': 'org-' + str(org.pk),
            'tipo': 'organismo',
            'nombre': org.nombre,
            'num_fuentes': fuentes_por_org.get(org.pk, 0),
        })
        if solo_relacionados:
            subs = org.suborganismos.filter(pk__in=sub_ids_con_insc)
        else:
            subs = org.suborganismos.all()
        for sub in subs:
            nodos.append({
                'id': 'sub-' + str(sub.pk),
                'tipo': 'suborganismo',
                'nombre': sub.nombre,
                'organismo_id': 'org-' + str(org.pk),
            })
            aristas.append({
                'origen': 'org-' + str(org.pk), 'destino': 'sub-' + str(sub.pk),
                'tipo': 'jerarquia', 'peso': 1,
            })

    # --- Equipos deduplicados por nombre normalizado + organismo ---
    equipo_groups = {}  # (nombre_norm, organismo_id) -> [insc, ...]
    for insc in inscripciones:
        nombre_equipo = insc.equipo_responsable or insc.nombre_completo or 'Equipo sin nombre'
        key = (_norm(nombre_equipo), insc.organismo_id)
        equipo_groups.setdefault(key, []).append(insc)

    insc_to_equipo = {}  # insc.pk -> equipo_id
    for (nombre_norm, org_id), group in equipo_groups.items():
        first = group[0]
        equipo_id = 'eq-' + str(first.pk)
        for insc in group:
            insc_to_equipo[insc.pk] = equipo_id

        # Suborganismo más frecuente del grupo (no siempre el del primer registro)
        sub_comun = Counter(i.suborganismo_id for i in group if i.suborganismo_id).most_common(1)
        sub_id = 'sub-' + str(sub_comun[0][0]) if sub_comun else None
        org_node_id = 'org-' + str(org_id) if org_id else None
        responsables = sorted(set(i.nombre_completo for i in group))
        total_fuentes = sum(i.fuentes.count() for i in group)

        nodos.append({
            'id': equipo_id,
            'tipo': 'equipo',
            'nombre': first.equipo_responsable or first.nombre_completo,
            'organismo': first.organismo.nombre if first.organismo else '',
            'organismo_id': org_node_id or '',
            'suborganismo_id': sub_id or '',
            'responsable': ', '.join(responsables),
            'num_fuentes': total_fuentes,
            'num_inscripciones': len(group),
        })
        if sub_id:
            aristas.append({
                'origen': sub_id, 'destino': equipo_id,
                'tipo': 'jerarquia', 'peso': 1,
            })
        elif org_node_id:
            aristas.append({
                'origen': org_node_id, 'destino': equipo_id,
                'tipo': 'pertenencia', 'peso': 1,
            })

    # --- Conjuntos (fuentes registradas) ---
    conjunto_vars_norm = {}  # fte_node_id -> set of normalized var names
    for fte in fuentes:
        org = fte.inscripcion.organismo
        insc = fte.inscripcion
        equipo_id = insc_to_equipo.get(insc.pk, 'eq-' + str(insc.pk))
        nodo_id = 'fte-' + str(fte.pk)
        vars_fte = list(fte.variables.all())
        vars_orig = [v.nombre for v in vars_fte]
        vars_norm = set(_norm_var(v) for v in vars_orig)
        conjunto_vars_norm[nodo_id] = vars_norm

        seguridad = getattr(fte, 'seguridad', None)
        nodos.append({
            'id': nodo_id,
            'tipo': 'conjunto',
            'nombre': fte.nombre,
            'fuente_id': fte.pk,
            'descripcion': fte.descripcion or '',
            'decisiones_apoya': fte.decisiones_apoya or '',
            'rol': fte.get_rol_display() if fte.rol else '',
            'criticidad': seguridad.criticidad if seguridad else '',
            'organismo': org.nombre if org else '',
            'organismo_id': 'org-' + str(org.pk) if org else '',
            'equipo_id': equipo_id,
            'num_variables': len(vars_orig),
            'num_identificadores': sum(1 for v in vars_fte if v.es_identificador),
            'variables': vars_orig,
        })
        aristas.append({
            'origen': equipo_id, 'destino': nodo_id,
            'tipo': 'pertenencia', 'peso': 1,
        })

        # Matching: variables del conjunto vs campos del catálogo (normalizado + fuzzy)
        bases_match = set()
        for vnorm in vars_norm:
            if not vnorm:
                continue
            # Exact match on normalized name
            if vnorm in campos_catalogo_norm:
                for bid in campos_catalogo_norm[vnorm]:
                    bases_match.add(bid)
            else:
                # Substring match (either direction, min 4 chars to avoid noise)
                for cnorm, bids in campos_catalogo_norm.items():
                    if len(vnorm) >= 4 and len(cnorm) >= 4 and (vnorm in cnorm or cnorm in vnorm):
                        for bid in bids:
                            bases_match.add(bid)
        for bid in bases_match:
            aristas.append({
                'origen': nodo_id, 'destino': bid,
                'tipo': 'matching_variable', 'peso': 1,
            })

    # --- Conectar conjuntos que comparten variables (>= 2 en común) ---
    fte_ids = list(conjunto_vars_norm.keys())
    for i in range(len(fte_ids)):
        for j in range(i + 1, len(fte_ids)):
            shared = conjunto_vars_norm[fte_ids[i]] & conjunto_vars_norm[fte_ids[j]]
            shared.discard('')
            if len(shared) >= 2:
                aristas.append({
                    'origen': fte_ids[i], 'destino': fte_ids[j],
                    'tipo': 'variables_compartidas', 'peso': len(shared),
                    'variables': sorted(shared),
                })

    # --- Fuentes externas ---
    fuentes_externas = {}  # nombre_norm -> {id, nombre, organismos: set}

    def obtener_o_crear_externa(nombre):
        key = _norm(nombre)
        if key not in fuentes_externas:
            fuentes_externas[key] = {
                'id': 'ext-' + str(len(fuentes_externas)),
                'nombre': nombre.strip(),
                'organismos': set(),
                'num_origenes': 0,
            }
        return fuentes_externas[key]

    # --- Índices de matching exacto (normalizado) ---
    catalogo_exacto = {}  # nombre_norm -> 'cat-<id>'
    catalogo_organismos = {}  # 'cat-<id>' -> organismo_normalizado
    for b in bases:
        bnorm = _norm(b['nombre'])
        if not bnorm:
            continue
        catalogo_exacto[bnorm] = 'cat-' + b['id']
        catalogo_organismos['cat-' + b['id']] = _norm(b.get('organismo', ''))

    # Índice de conjuntos pre-normalizado (una sola pasada, no por referencia)
    conjunto_exacto = {}  # nombre_norm -> 'fte-<pk>'
    fuentes_norm = []  # (pk, nombre_norm, organismo_norm)
    for fte in fuentes:
        fnorm = _norm(fte.nombre)
        if not fnorm:
            continue
        conjunto_exacto.setdefault(fnorm, 'fte-' + str(fte.pk))
        forg = _norm(fte.inscripcion.organismo.nombre) if fte.inscripcion.organismo else ''
        fuentes_norm.append((fte.pk, fnorm, forg))

    def _overlap(a, b):
        """Proporción de solape entre dos cadenas normalizadas (0-1)."""
        return min(len(a), len(b)) / max(len(a), len(b))

    def match_catalogo(nombre_norm, organismo_norm=''):
        if nombre_norm in catalogo_exacto:
            return catalogo_exacto[nombre_norm], 'match_exacto'
        if len(nombre_norm) < 4:
            return None, None
        candidatos = []
        for bnorm, cat_id in catalogo_exacto.items():
            if len(bnorm) >= 4 and (nombre_norm in bnorm or bnorm in nombre_norm):
                candidatos.append((_overlap(nombre_norm, bnorm), cat_id))
        if not candidatos:
            return None, None
        candidatos.sort(key=lambda c: c[0], reverse=True)
        if organismo_norm:
            for _, cat_id in candidatos:
                cat_org = catalogo_organismos.get(cat_id, '')
                if cat_org and (organismo_norm in cat_org or cat_org in organismo_norm):
                    return cat_id, 'match_exacto'
        return candidatos[0][1], 'match_aproximado'

    def match_conjunto(nombre_norm, organismo_norm='', excluir_pk=None):
        exacto = conjunto_exacto.get(nombre_norm)
        if exacto and exacto != 'fte-' + str(excluir_pk):
            return exacto, 'match_exacto'
        if len(nombre_norm) < 4:
            return None, None
        candidatos = []
        for pk, fnorm, forg in fuentes_norm:
            if pk == excluir_pk:
                continue
            if len(fnorm) >= 4 and (nombre_norm in fnorm or fnorm in nombre_norm):
                candidatos.append((_overlap(nombre_norm, fnorm), pk, forg))
        if not candidatos:
            return None, None
        candidatos.sort(key=lambda c: c[0], reverse=True)
        if organismo_norm:
            for _, pk, forg in candidatos:
                if forg and (organismo_norm in forg or forg in organismo_norm):
                    return 'fte-' + str(pk), 'match_exacto'
        return 'fte-' + str(candidatos[0][1]), 'match_aproximado'

    # --- Aristas: FuenteOrigen ---
    for fte in fuentes:
        org = fte.inscripcion.organismo
        for orig in fte.fuentes_origen.all():
            nombre_norm = _norm(orig.nombre)
            if not nombre_norm:
                continue
            org_norm = _norm(orig.organismo_productor) if orig.organismo_productor else ''
            destino, tipo_match = match_catalogo(nombre_norm, org_norm)
            if destino:
                aristas.append({
                    'origen': 'fte-' + str(fte.pk), 'destino': destino,
                    'tipo': tipo_match, 'peso': 3 if tipo_match == 'match_exacto' else 1,
                })
            else:
                destino, tipo_match = match_conjunto(nombre_norm, org_norm, excluir_pk=fte.pk)
                if destino:
                    aristas.append({
                        'origen': 'fte-' + str(fte.pk), 'destino': destino,
                        'tipo': tipo_match, 'peso': 3 if tipo_match == 'match_exacto' else 1,
                    })
                else:
                    ext = obtener_o_crear_externa(orig.nombre)
                    ext['organismos'].add(org.nombre if org else '')
                    ext['num_origenes'] += 1
                    aristas.append({
                        'origen': 'fte-' + str(fte.pk), 'destino': ext['id'],
                        'tipo': 'origen_externo', 'peso': 1,
                    })

    # --- Retos: nodos + aristas FuenteRequerida ---
    retos = (RetoAnalisis.objects
             .select_related('inscripcion__organismo')
             .prefetch_related('fuentes_requeridas').all())
    for reto in retos:
        org = reto.inscripcion.organismo if reto.inscripcion else None
        reto_id = 'reto-' + str(reto.pk)
        equipo_id = insc_to_equipo.get(reto.inscripcion_id, '') if reto.inscripcion_id else ''

        nodos.append({
            'id': reto_id,
            'tipo': 'reto',
            'nombre': reto.nombre or 'Reto sin nombre',
            'descripcion': (reto.analisis_propuesto or '')[:200],
            'organismo': org.nombre if org else '',
            'organismo_id': 'org-' + str(org.pk) if org else '',
            'equipo_id': equipo_id,
        })
        if equipo_id:
            aristas.append({
                'origen': equipo_id, 'destino': reto_id,
                'tipo': 'pertenencia', 'peso': 1,
            })

        for req in reto.fuentes_requeridas.all():
            nombre_norm = _norm(req.nombre)
            if not nombre_norm:
                continue
            org_norm = _norm(req.organismo) if req.organismo else ''
            destino, tipo_match = match_catalogo(nombre_norm, org_norm)
            if not destino:
                destino, tipo_match = match_conjunto(nombre_norm, org_norm)
            if destino:
                aristas.append({
                    'origen': reto_id, 'destino': destino,
                    'tipo': tipo_match, 'peso': 3 if tipo_match == 'match_exacto' else 1,
                })
            else:
                ext = obtener_o_crear_externa(req.nombre)
                ext['organismos'].add(org.nombre if org else '')
                ext['num_origenes'] += 1
                aristas.append({
                    'origen': reto_id, 'destino': ext['id'],
                    'tipo': 'requerida_externa', 'peso': 1,
                })

    # --- Nodos de fuentes externas ---
    for key, ext in fuentes_externas.items():
        nodos.append({
            'id': ext['id'],
            'tipo': 'fuente_informacion',
            'nombre': ext['nombre'],
            'organismos': ', '.join(sorted(ext['organismos'])) or '—',
            'num_origenes': ext['num_origenes'],
        })

    # Deduplicar aristas idénticas (mismo origen, destino y tipo)
    vistas = set()
    aristas_unicas = []
    for a in aristas:
        clave = (a['origen'], a['destino'], a['tipo'])
        if clave in vistas:
            continue
        vistas.add(clave)
        aristas_unicas.append(a)

    if solo_relacionados:
        # Conservar solo nodos con al menos una arista de RELACIÓN (no
        # estructural), más sus ancestros para no romper la jerarquía.
        TIPOS_RELACION = {'match_exacto', 'match_aproximado', 'origen_externo',
                          'requerida_externa', 'matching_variable',
                          'variables_compartidas'}
        con_relacion = set()
        for a in aristas_unicas:
            if a['tipo'] in TIPOS_RELACION:
                con_relacion.add(a['origen'])
                con_relacion.add(a['destino'])
        por_id = {n['id']: n for n in nodos}
        conservar = set(con_relacion)
        pila = list(con_relacion)
        while pila:
            nodo = por_id.get(pila.pop())
            if not nodo:
                continue
            for clave_padre in ('equipo_id', 'suborganismo_id', 'organismo_id'):
                pid = nodo.get(clave_padre)
                if pid and pid not in conservar:
                    conservar.add(pid)
                    pila.append(pid)
        # El catálogo (incluye Gold) siempre se conserva: es la capa de referencia
        nodos = [n for n in nodos
                 if n['id'] in conservar or n['tipo'] in ('gold', 'catalogo')]
        ids_nodos = {n['id'] for n in nodos}
        aristas_unicas = [a for a in aristas_unicas
                          if a['origen'] in ids_nodos and a['destino'] in ids_nodos]

    return nodos, aristas_unicas


def _grafo_ecosistema(solo_relacionados=False):
    """Versión cacheada del constructor del grafo (30 s): los polls del 3D y
    los exports comparten el mismo build en lugar de recalcularlo por request."""
    from django.core.cache import cache
    clave = 'grafo_ecosistema_rel' if solo_relacionados else 'grafo_ecosistema'
    grafo = cache.get(clave)
    if grafo is None:
        grafo = _construir_grafo_ecosistema(solo_relacionados)
        cache.set(clave, grafo, 30)
    return grafo


@require_GET
def api_ecosistema(request):
    """Endpoint JSON con nodos y aristas para la visualización 3D del ecosistema."""
    solo_relacionados = request.GET.get('solo_relacionados', 'false') == 'true'
    nodos, aristas = _grafo_ecosistema(solo_relacionados)
    return JsonResponse({'nodos': nodos, 'aristas': aristas})


@require_GET
def ecosistema_gexf(request):
    """Descarga el grafo del ecosistema en formato GEXF (Gephi / Gephi Lite)."""
    from xml.sax.saxutils import escape
    from django.http import HttpResponse

    solo_relacionados = request.GET.get('solo_relacionados', 'false') == 'true'
    nodos, aristas = _grafo_ecosistema(solo_relacionados)
    # Posiciones pre-calculadas: Gephi/Gephi Lite abren el grafo ya dispuesto
    # con el mismo layout concéntrico que el SVG renderizado
    pos = _layout_anillos(nodos)

    def _esc(s):
        return escape(str(s), {'"': '&quot;'})

    def _rgb(hex_color):
        h = hex_color.lstrip('#')
        return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)

    partes = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<gexf xmlns="http://www.gexf.net/1.3" xmlns:viz="http://www.gexf.net/1.3/viz" version="1.3">',
        '<meta><creator>Ecosistema de Datos</creator>'
        '<description>Mapa del ecosistema de datos institucionales</description></meta>',
        '<graph mode="static" defaultedgetype="undirected">',
        '<attributes class="node">',
        '  <attribute id="tipo" title="Tipo" type="string"/>',
        '  <attribute id="descripcion" title="Descripcion" type="string"/>',
        '  <attribute id="num_campos" title="Num Campos" type="integer"/>',
        '  <attribute id="organismo" title="Organismo" type="string"/>',
        '  <attribute id="organismos" title="Organismos" type="string"/>',
        '  <attribute id="num_origenes" title="Num Origenes" type="integer"/>',
        '  <attribute id="criticidad" title="Criticidad" type="string"/>',
        '  <attribute id="rol" title="Rol" type="string"/>',
        '  <attribute id="num_variables" title="Num Variables" type="integer"/>',
        '  <attribute id="responsable" title="Responsable" type="string"/>',
        '  <attribute id="num_fuentes" title="Num Fuentes" type="integer"/>',
        '  <attribute id="num_inscripciones" title="Num Inscripciones" type="integer"/>',
        '</attributes>',
        '<attributes class="edge">',
        '  <attribute id="tipo" title="Tipo" type="string"/>',
        '  <attribute id="peso" title="Peso" type="float"/>',
        '  <attribute id="variables" title="Variables" type="string"/>',
        '  <attribute id="campo" title="Campo" type="string"/>',
        '</attributes>',
        '<nodes>',
    ]
    for n in nodos:
        r, g, b = _rgb(COLOR_NODO.get(n.get('tipo'), '#888888'))
        size = max(5, _tamano_nodo(n) * 2)
        partes.append(f'  <node id="{_esc(n["id"])}" label="{_esc(n.get("nombre", ""))}">')
        partes.append('    <attvalues>')
        partes.append(f'      <attvalue for="tipo" value="{_esc(TIPO_LABEL.get(n.get("tipo"), n.get("tipo", "")))}"/>')
        if n.get('descripcion'):
            partes.append(f'      <attvalue for="descripcion" value="{_esc(n["descripcion"])}"/>')
        if n.get('num_campos') is not None:
            partes.append(f'      <attvalue for="num_campos" value="{n["num_campos"]}"/>')
        if n.get('organismo'):
            partes.append(f'      <attvalue for="organismo" value="{_esc(n["organismo"])}"/>')
        if n.get('organismos'):
            partes.append(f'      <attvalue for="organismos" value="{_esc(n["organismos"])}"/>')
        if n.get('num_origenes') is not None:
            partes.append(f'      <attvalue for="num_origenes" value="{n["num_origenes"]}"/>')
        if n.get('criticidad'):
            partes.append(f'      <attvalue for="criticidad" value="{_esc(n["criticidad"])}"/>')
        if n.get('rol'):
            partes.append(f'      <attvalue for="rol" value="{_esc(n["rol"])}"/>')
        if n.get('num_variables') is not None:
            partes.append(f'      <attvalue for="num_variables" value="{n["num_variables"]}"/>')
        if n.get('responsable'):
            partes.append(f'      <attvalue for="responsable" value="{_esc(n["responsable"])}"/>')
        if n.get('num_fuentes') is not None:
            partes.append(f'      <attvalue for="num_fuentes" value="{n["num_fuentes"]}"/>')
        if n.get('num_inscripciones') is not None:
            partes.append(f'      <attvalue for="num_inscripciones" value="{n["num_inscripciones"]}"/>')
        partes.append('    </attvalues>')
        partes.append(f'    <viz:color r="{r}" g="{g}" b="{b}" a="0.9"/>')
        partes.append(f'    <viz:size value="{size:.2f}"/>')
        p = pos.get(n['id'])
        if p:
            partes.append(f'    <viz:position x="{p[0]:.2f}" y="{p[1]:.2f}" z="0"/>')
        partes.append('  </node>')
    partes.append('</nodes>')
    partes.append('<edges>')
    for i, a in enumerate(aristas):
        r, g, b = _rgb(COLOR_ARISTA.get(a.get('tipo'), '#666666'))
        peso = a.get('peso') or 1
        partes.append(
            f'  <edge id="e{i}" source="{_esc(a["origen"])}" '
            f'target="{_esc(a["destino"])}" weight="{peso}">')
        partes.append('    <attvalues>')
        partes.append(f'      <attvalue for="tipo" value="{_esc(a.get("tipo", ""))}"/>')
        partes.append(f'      <attvalue for="peso" value="{peso}"/>')
        if a.get('variables'):
            vars_str = ', '.join(a['variables'])
            partes.append(f'      <attvalue for="variables" value="{_esc(vars_str)}"/>')
        if a.get('campo'):
            partes.append(f'      <attvalue for="campo" value="{_esc(a["campo"])}"/>')
        partes.append('    </attvalues>')
        partes.append(f'    <viz:color r="{r}" g="{g}" b="{b}" a="0.4"/>')
        partes.append(f'    <viz:thickness value="{peso * 0.5 + 0.5:.2f}"/>')
        partes.append('  </edge>')
    partes.append('</edges>')
    partes.append('</graph>\n</gexf>')

    response = HttpResponse('\n'.join(partes), content_type='application/gexf+xml; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="ecosistema_datos.gexf"'
    return response


def _layout_anillos(nodos):
    """Layout determinista de anillos concéntricos por tipo de nodo (misma
    metáfora que la vista 3D): Gold al centro y cada capa semántica en su
    anillo, con los hijos agrupados cerca del ángulo del padre.
    Compartido por el SVG renderizado y las posiciones del GEXF."""
    import math

    por_tipo = {}
    for n in nodos:
        por_tipo.setdefault(n.get('tipo'), []).append(n)
    pos = {}

    def _angulo_padre(n, claves):
        for k in claves:
            pid = n.get(k)
            if pid and pid in pos:
                px, py = pos[pid]
                return math.atan2(py, px)
        return 0.0

    def anillo(lista, radio, claves_padre=()):
        if not lista:
            return
        if claves_padre:
            lista = sorted(lista, key=lambda n: _angulo_padre(n, claves_padre))
        # El radio crece si el anillo no cabe (arco mínimo ~26px por nodo)
        radio = max(radio, len(lista) * 26 / (2 * math.pi))
        for i, n in enumerate(lista):
            ang = (i / len(lista)) * 2 * math.pi - math.pi / 2
            pos[n['id']] = (math.cos(ang) * radio, math.sin(ang) * radio)

    anillo(por_tipo.get('gold', []), 0)
    anillo(por_tipo.get('catalogo', []), 300)
    anillo(por_tipo.get('organismo', []), 640)
    anillo(por_tipo.get('suborganismo', []), 980, ('organismo_id',))
    anillo(por_tipo.get('equipo', []), 1320, ('suborganismo_id', 'organismo_id'))
    anillo(por_tipo.get('conjunto', []) + por_tipo.get('reto', []), 1660, ('equipo_id',))
    anillo(por_tipo.get('fuente_informacion', []), 2000)
    return pos


@require_GET
def ecosistema_svg(request):
    """Renderiza el grafo del ecosistema como imagen SVG con layout de anillos
    concéntricos por tipo (misma metáfora que la vista 3D y las posiciones del
    GEXF). Acepta ?solo_relacionados=true. Descarga el archivo (imprimible a PDF)."""
    import math
    from xml.sax.saxutils import escape
    from django.http import HttpResponse

    solo_relacionados = request.GET.get('solo_relacionados', 'false') == 'true'
    nodos, aristas = _grafo_ecosistema(solo_relacionados)

    def _esc(s):
        return escape(str(s))

    pos = _layout_anillos(nodos)

    radio_max = max((math.hypot(x, y) for x, y in pos.values()), default=0)
    margen = 300  # espacio para las etiquetas radiales del anillo externo
    lado = int((radio_max + margen) * 2)
    c = lado / 2  # centro del canvas

    partes = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{lado}" height="{lado}" '
        f'viewBox="0 0 {lado} {lado}" font-family="Montserrat, Arial, sans-serif">',
        f'<rect width="{lado}" height="{lado}" fill="#0b1020"/>',
    ]

    # Guías de anillos: hace visible la estructura concéntrica
    radios = sorted({round(math.hypot(x, y)) for x, y in pos.values()
                     if math.hypot(x, y) > 1})
    for rad in radios:
        partes.append(
            f'<circle cx="{c}" cy="{c}" r="{rad}" fill="none" stroke="#ffffff" '
            f'stroke-opacity="0.07" stroke-dasharray="3 7"/>')

    # Aristas (debajo de los nodos)
    for a in aristas:
        p1, p2 = pos.get(a['origen']), pos.get(a['destino'])
        if not p1 or not p2:
            continue
        color = COLOR_ARISTA.get(a.get('tipo'), '#666666')
        partes.append(
            f'<line x1="{p1[0] + c:.1f}" y1="{p1[1] + c:.1f}" '
            f'x2="{p2[0] + c:.1f}" y2="{p2[1] + c:.1f}" '
            f'stroke="{color}" stroke-width="1" stroke-opacity="0.35">'
            f'<title>{_esc(a.get("tipo", ""))}</title></line>')

    # Nodos (tooltip nativo vía <title>) + etiqueta radial en TODOS los nodos:
    # el texto sigue la dirección del anillo hacia afuera, como un dendrograma
    for n in nodos:
        p = pos.get(n['id'])
        if not p:
            continue
        tipo = n.get('tipo')
        color = COLOR_NODO.get(tipo, '#888888')
        r = max(3, _tamano_nodo(n))
        label = f'{TIPO_LABEL.get(tipo, tipo)}: {n.get("nombre", "")}'
        partes.append(
            f'<circle cx="{p[0] + c:.1f}" cy="{p[1] + c:.1f}" r="{r:.1f}" '
            f'fill="{color}" fill-opacity="0.92" stroke="#ffffff" '
            f'stroke-opacity="0.25" stroke-width="0.6">'
            f'<title>{_esc(label)}</title></circle>')
        nombre = _esc((n.get('nombre') or '')[:34])
        if not nombre:
            continue
        if tipo == 'gold':
            partes.append(
                f'<text x="{p[0] + c:.1f}" y="{p[1] + c + r + 18:.1f}" '
                f'font-size="15" font-weight="bold" fill="#ffd700" '
                f'text-anchor="middle" '
                f'style="paint-order:stroke;stroke:#0b1020;stroke-width:3">'
                f'{nombre}</text>')
            continue
        ang = math.atan2(p[1], p[0])
        deg = math.degrees(ang)
        lx = p[0] + math.cos(ang) * (r + 5)
        ly = p[1] + math.sin(ang) * (r + 5)
        # Mitad izquierda: rotar 180° y anclar al final para leer hacia afuera
        rot, anchor = (deg + 180, 'end') if (deg > 90 or deg < -90) else (deg, 'start')
        fs = 10 if tipo in ('catalogo', 'organismo') else 8.5
        partes.append(
            f'<text x="{lx + c:.1f}" y="{ly + c:.1f}" font-size="{fs}" '
            f'fill="#cbd5e0" text-anchor="{anchor}" dominant-baseline="middle" '
            f'transform="rotate({rot:.1f} {lx + c:.1f} {ly + c:.1f})" '
            f'style="paint-order:stroke;stroke:#0b1020;stroke-width:2.5">'
            f'{nombre}</text>')

    # Leyenda
    lx, ly = 24, 30
    partes.append(f'<g font-size="13" fill="#e2e8f0">')
    partes.append(f'<text x="{lx}" y="{ly - 10}" font-size="15" font-weight="bold">Ecosistema de datos</text>')
    for i, (tipo, label) in enumerate(TIPO_LABEL.items()):
        y = ly + 14 + i * 20
        partes.append(f'<circle cx="{lx + 6}" cy="{y - 4}" r="5" fill="{COLOR_NODO[tipo]}"/>')
        partes.append(f'<text x="{lx + 18}" y="{y}">{_esc(label)}</text>')
    partes.append('</g>')

    partes.append('</svg>')
    response = HttpResponse('\n'.join(partes), content_type='image/svg+xml; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="ecosistema_datos.svg"'
    return response


def _inscripcion_activa(request):
    """Inscripción activa del usuario: la de la sesión si existe, si no la
    primera registrada (un usuario puede tener varias vía admin)."""
    if not request.user.is_authenticated:
        return None
    qs = Inscripcion.objects.filter(user=request.user)
    insc_id = request.session.get(SESSION_KEY)
    if insc_id:
        insc = qs.filter(pk=insc_id).first()
        if insc:
            return insc
    return qs.first()


@delegado_required
def wizard_inicio(request):
    inscripcion = _inscripcion_activa(request)
    if request.method == 'POST':
        if inscripcion and request.POST.get('accion') == 'continuar':
            # Si la inscripción está completada, ir al paso de fuentes
            # para permitir agregar o editar registros de fuentes de datos
            if inscripcion.estado == Inscripcion.Estado.COMPLETADA:
                return redirect('indicadores:taller_wizard_paso', slug='fuente-general')
            destino = inscripcion.paso_actual or wz.PASOS[0].slug
            if destino == 'identificacion' and inscripcion.ficha_inicial == 'f2':
                siguiente = wz.PASOS_POR_SLUG.get('reto-definicion')
            elif destino == 'reto-tratamiento' and inscripcion.ficha_inicial == 'f2':
                return redirect('indicadores:taller_wizard_confirmacion')
            else:
                siguiente = wz.paso_siguiente(destino)
            return redirect('indicadores:taller_wizard_paso', slug=siguiente.slug if siguiente else destino)
        return redirect('indicadores:taller_wizard_paso', slug=wz.PASOS[0].slug)
    context = {
        'inscripcion': inscripcion,
        'pasos': wz.PASOS,
        'fases': wz.FASES,
    }
    return render(request, 'taller/wizard_inicio.html', context)


@delegado_required
def wizard_paso(request, slug):
    paso = wz.PASOS_POR_SLUG.get(slug)
    if paso is None:
        raise Http404

    inscripcion = _inscripcion_activa(request)
    es_primer_paso = paso.slug == wz.PASOS[0].slug
    if inscripcion is None and not es_primer_paso:
        return redirect('indicadores:taller_wizard')

    # --- Determinar la instancia base ---
    if paso.slug == 'fuente-general':
        # El formset de múltiples fuentes usa la inscripción como instancia
        instancia = inscripcion
    elif paso.slug in PASOS_CON_FUENTE:
        # Seleccionar la fuente activa de la sesión o la primera disponible
        fuente_id = request.session.get(FUENTE_KEY)
        fuentes_qs = FuenteDatos.objects.filter(inscripcion=inscripcion)
        fuente_sel = fuentes_qs.filter(pk=fuente_id).first() if fuente_id else None
        if fuente_sel is None:
            fuente_sel = fuentes_qs.first()
        if fuente_sel is None:
            # Si no hay fuentes creadas, redirigir al paso 2
            return redirect('indicadores:taller_wizard_paso', slug='fuente-general')
        if paso.slug in PASOS_SEGURIDAD:
            instancia, _ = EvaluacionSeguridad.objects.get_or_create(fuente=fuente_sel)
        else:
            instancia = fuente_sel
    else:
        fuente_sel = None
        instancia = paso.get_instance(inscripcion) if inscripcion else None

    # --- Lista de fuentes para el selector (pasos F1 individuales y F3) ---
    fuentes_lista = None
    if paso.slug in PASOS_CON_FUENTE and inscripcion:
        fuentes_lista = FuenteDatos.objects.filter(inscripcion=inscripcion).order_by('pk')
        if fuente_sel:
            request.session[FUENTE_KEY] = fuente_sel.pk

    form = None
    formset = None
    if request.method == 'POST':
        # Cambio de fuente activa sin guardar
        if paso.slug in PASOS_CON_FUENTE and request.POST.get('cambiar_fuente'):
            try:
                request.session[FUENTE_KEY] = int(request.POST.get('fuente_activa'))
            except (TypeError, ValueError):
                pass  # valor malformado: se ignora y se recarga el paso
            return redirect('indicadores:taller_wizard_paso', slug=slug)

        # En el paso de rol, las fuentes de origen solo aplican para Usuario/Ambos
        omitir_formset = paso.slug == 'fuente-rol' and request.POST.get('rol') not in ('usuario', 'ambos')
        if paso.form_class:
            form = paso.form_class(request.POST, request.FILES, instance=instancia)
        if paso.formset_class and not omitir_formset:
            formset = paso.formset_class(request.POST, instance=instancia)
        form_ok = form.is_valid() if form else True
        formset_ok = formset.is_valid() if formset else True
        if form_ok and formset_ok:
            if form:
                objeto = form.save(commit=False)
                if es_primer_paso and inscripcion is None:
                    objeto.user = request.user
                objeto.save()
                # save(commit=False) no persiste los M2M: hay que llamarlo
                # explícitamente o los checkboxes (problemas_calidad,
                # tratamientos, etc.) se pierden silenciosamente.
                if hasattr(form, 'save_m2m'):
                    form.save_m2m()
                if es_primer_paso and inscripcion is None:
                    inscripcion = objeto
                    request.session[SESSION_KEY] = inscripcion.pk
            if formset:
                formset.save()
                # Si se crearon fuentes nuevas, activar la última para que los
                # pasos siguientes (actualización, rol, variables...) la editen
                if paso.slug == 'fuente-general' and getattr(formset, 'new_objects', None):
                    request.session[FUENTE_KEY] = formset.new_objects[-1].pk
            if omitir_formset and instancia is not None:
                instancia.fuentes_origen.all().delete()
            if wz.indice(paso.slug) >= wz.indice(inscripcion.paso_actual):
                inscripcion.paso_actual = paso.slug
            inscripcion.save(update_fields=['paso_actual', 'actualizado'])
            if paso.slug == 'identificacion' and inscripcion.ficha_inicial == 'f2':
                siguiente = wz.PASOS_POR_SLUG.get('reto-definicion')
            elif paso.slug == 'reto-tratamiento' and inscripcion.ficha_inicial == 'f2':
                siguiente = None
            else:
                siguiente = wz.paso_siguiente(paso.slug)
            if siguiente:
                return redirect('indicadores:taller_wizard_paso', slug=siguiente.slug)
            return redirect('indicadores:taller_wizard_confirmacion')
    else:
        if paso.form_class:
            form = paso.form_class(instance=instancia)
        if paso.formset_class:
            formset = paso.formset_class(instance=instancia)

    numero = wz.indice(paso.slug) + 1
    context = {
        'paso': paso,
        'form': form,
        'formset': formset,
        'fase': wz.FASES[paso.fase],
        'fases': wz.FASES,
        'pasos': wz.PASOS,
        'numero': numero,
        'total': len(wz.PASOS),
        'progreso': int(numero * 100 / (len(wz.PASOS) + 1)),
        'anterior': wz.PASOS_POR_SLUG.get('identificacion') if inscripcion and paso.slug == 'reto-definicion' and inscripcion.ficha_inicial == 'f2' else wz.paso_anterior(paso.slug),
        'inscripcion': inscripcion,
        'fuentes_lista': fuentes_lista,
        'fuente_activa': fuente_sel if paso.slug in PASOS_CON_FUENTE else None,
    }
    if paso.slug == 'seguridad-valoracion' and instancia and instancia.pk:
        context['ocha'] = instancia.calcular_factores_riesgo_ocha()
    return render(request, 'taller/wizard_paso.html', context)


@delegado_required
def wizard_confirmacion(request):
    inscripcion = _inscripcion_activa(request)
    if inscripcion is None:
        return redirect('indicadores:taller_wizard')
    if request.method == 'POST':
        inscripcion.estado = Inscripcion.Estado.COMPLETADA
        inscripcion.save(update_fields=['estado', 'actualizado'])
        request.session.pop(SESSION_KEY, None)
        return render(request, 'taller/wizard_exito.html', {'inscripcion': inscripcion})
    context = {
        'inscripcion': inscripcion,
        'pasos': wz.PASOS,
        'fases': wz.FASES,
    }
    # OCHA de la fuente activa (o la primera) — F3 es por fuente
    fuente_id = request.session.get(FUENTE_KEY)
    fuente_conf = FuenteDatos.objects.filter(inscripcion=inscripcion)
    fuente_conf = (fuente_conf.filter(pk=fuente_id).first() if fuente_id else None) or fuente_conf.first()
    seguridad = getattr(fuente_conf, 'seguridad', None) if fuente_conf else None
    if seguridad and seguridad.pk:
        context['ocha'] = seguridad.calcular_factores_riesgo_ocha()
    return render(request, 'taller/wizard_confirmacion.html', context)


SEMAFORO_COLORES = {'verde': '#2f9e57', 'amarillo': '#f0b429', 'rojo': '#e53e3e', '': '#a0aec0'}


def _filas(pares):
    """Convierte pares (etiqueta, total) en filas con % relativo al máximo para las barras."""
    pares = [p for p in pares if p[1] > 0]
    maximo = max((total for _, total in pares), default=0)
    return [{
        'label': etiqueta,
        'total': total,
        'pct': int(total * 100 / maximo) if maximo else 0,
        'color': None,
    } for etiqueta, total in pares]


def _panel_m2m(titulo, m2m, ids):
    campo_padre = [f.name for f in m2m.through._meta.fields
                   if f.is_relation and f.related_model is not Inscripcion and f.name != 'opciontaller'][0]
    modelo_padre = m2m.through._meta.get_field(campo_padre).related_model
    tiene_insc = any(f.name == 'inscripcion' for f in modelo_padre._meta.fields)
    ruta = (f'{campo_padre}__inscripcion_id__in' if tiene_insc
            else f'{campo_padre}__fuente__inscripcion_id__in')
    datos = (m2m.through.objects
             .filter(**{ruta: ids})
             .values('opciontaller__etiqueta')
             .annotate(total=Count('id'))
             .order_by('-total'))
    return {'titulo': titulo,
            'filas': _filas([(d['opciontaller__etiqueta'], d['total']) for d in datos])}


def _panel_choice(titulo, qs, campo, semaforo=False, incluir_vacios=True):
    modelo = qs.model
    etiquetas = dict(modelo._meta.get_field(campo).choices)
    datos = qs.values(campo).annotate(total=Count('id')).order_by('-total')
    pares = []
    for d in datos:
        clave = d[campo]
        if clave == '' and not incluir_vacios:
            continue
        etiqueta = etiquetas.get(clave, 'Sin responder') if clave else 'Sin responder'
        pares.append((etiqueta, d['total'], clave))
    filas = []
    maximo = max((t for _, t, _ in pares), default=0)
    for etiqueta, total, clave in pares:
        filas.append({
            'label': etiqueta,
            'total': total,
            'pct': int(total * 100 / maximo) if maximo else 0,
            'color': SEMAFORO_COLORES.get(clave) if semaforo else None,
        })
    return {'titulo': titulo, 'filas': filas}


def _panel_texto(titulo, qs, campo):
    """Agrega un TextField con valores separados por comas en un panel de barras."""
    from collections import Counter
    contador = Counter()
    for obj in qs:
        valor = getattr(obj, campo, '')
        if valor:
            for v in valor.split(','):
                v = v.strip()
                if v:
                    contador[v] += 1
    pares = contador.most_common()
    return {'titulo': titulo, 'filas': _filas(pares)}


def _fuentes_detalle(qs, ids):
    """Construye resumen F1 de cada fuente para las tarjetas interactivas del dashboard."""
    fuentes = (qs.exclude(nombre='')
               .select_related('inscripcion__organismo')
               .prefetch_related('relaciones', 'instrumentos', 'mecanismos',
                                 'periodicidades', 'desagregaciones', 'unidades_analisis',
                                 'problemas_calidad', 'fuentes_origen', 'variables')
               .order_by('inscripcion__organismo__nombre', 'nombre'))
    resultado = []
    for f in fuentes:
        vars_qs = f.variables.all()
        resultado.append({
            'id': f.pk,
            'nombre': f.nombre,
            'organismo': f.inscripcion.organismo.nombre if f.inscripcion.organismo else '—',
            'descripcion': f.descripcion or 'Sin descripción',
            'decisiones_apoya': f.decisiones_apoya or '—',
            'rol': f.get_rol_display() if f.rol else '—',
            'relaciones': ', '.join(r.etiqueta for r in f.relaciones.all()) or '—',
            'instrumentos': ', '.join(i.etiqueta for i in f.instrumentos.all()) or '—',
            'mecanismos': ', '.join(m.etiqueta for m in f.mecanismos.all()) or '—',
            'periodicidades': ', '.join(p.etiqueta for p in f.periodicidades.all()) or '—',
            'desagregaciones': ', '.join(d.etiqueta for d in f.desagregaciones.all()) or '—',
            'fecha_actualizacion': f.fecha_actualizacion.strftime('%Y-%m-%d') if f.fecha_actualizacion else '—',
            'num_registros': f.num_registros if f.num_registros else ('Desconocido' if f.num_registros_desconocido else '—'),
            'georref': f.get_georreferenciacion_display() if f.georreferenciacion else '—',
            'tiene_diccionario': f.get_tiene_diccionario_display() if f.tiene_diccionario else '—',
            'tiene_validacion': f.get_tiene_validacion_display() if f.tiene_validacion else '—',
            'num_variables': vars_qs.count(),
            'num_identificadores': vars_qs.filter(es_identificador=True).count(),
            'num_origenes': f.fuentes_origen.count(),
            'unidades': ', '.join(u.etiqueta for u in f.unidades_analisis.all()) or '—',
            'problemas': ', '.join(p.etiqueta for p in f.problemas_calidad.all()) or '—',
        })
    return resultado


@taller_staff_required
def dashboard(request):
    from .models_taller import FuenteRequerida
    from .models_catalogo import EvaluacionSeguridad, VariableFuente
    from . import wizard_taller as wz

    alcance = request.GET.get('alcance', 'todas')
    inscripciones = Inscripcion.objects.select_related(
        'organismo', 'suborganismo', 'reto').all()
    if alcance == 'completadas':
        inscripciones = inscripciones.filter(estado=Inscripcion.Estado.COMPLETADA)
    ids = list(inscripciones.values_list('pk', flat=True))

    fuentes = FuenteDatos.objects.filter(inscripcion_id__in=ids)
    retos = RetoAnalisis.objects.filter(inscripcion_id__in=ids)
    seguridades = EvaluacionSeguridad.objects.filter(fuente__inscripcion_id__in=ids)
    variables = VariableFuente.objects.filter(fuente__inscripcion_id__in=ids)
    fuentes_requeridas = FuenteRequerida.objects.filter(reto__inscripcion_id__in=ids)

    todas = Inscripcion.objects.all()
    completadas = todas.filter(estado=Inscripcion.Estado.COMPLETADA)

    # Embudo: fichas que completaron cada paso
    pasos_alcanzados = list(inscripciones.values_list('paso_actual', 'estado'))
    embudo = []
    for i, p in enumerate(wz.PASOS):
        n = sum(1 for slug, _ in pasos_alcanzados if wz.indice(slug) >= i)
        embudo.append({'titulo': f'{i + 1}. {p.titulo}', 'fase': p.fase, 'total': n,
                       'pct': int(n * 100 / len(pasos_alcanzados)) if pasos_alcanzados else 0})

    # Organismos y fuentes más citadas
    organismos = _filas(list(inscripciones.values_list('organismo__nombre')
                             .annotate(total=Count('id')).order_by('-total')))
    fuentes_top = _filas(list(fuentes.exclude(nombre='').values_list('nombre')
                              .annotate(total=Count('id')).order_by('-total')))
    requeridas_top = _filas(list(fuentes_requeridas.exclude(nombre='').values_list('nombre')
                                 .annotate(total=Count('id')).order_by('-total')))

    paneles_f1 = [
        _panel_m2m('Relación con el conjunto de datos', FuenteDatos.relaciones, ids),
        _panel_m2m('Instrumento de recolección', FuenteDatos.instrumentos, ids),
        _panel_m2m('Mecanismo de consolidación', FuenteDatos.mecanismos, ids),
        _panel_m2m('Periodicidad de actualización', FuenteDatos.periodicidades, ids),
        _panel_m2m('Desagregación geográfica', FuenteDatos.desagregaciones, ids),
        _panel_m2m('Unidad de análisis principal', FuenteDatos.unidades_analisis, ids),
        _panel_choice('Rol del organismo frente a la fuente', fuentes, 'rol'),
        _panel_choice('¿Existe georreferenciación?', fuentes, 'georreferenciacion'),
        _panel_m2m('Nivel de georreferenciación', FuenteDatos.niveles_georref, ids),
        _panel_m2m('Problemas de calidad identificados', FuenteDatos.problemas_calidad, ids),
        _panel_choice('¿Procesos de validación o limpieza?', fuentes, 'tiene_validacion'),
        _panel_choice('¿Existe diccionario de variables?', fuentes, 'tiene_diccionario'),
    ]

    paneles_f2 = [
        _panel_m2m('Unidades que se desean relacionar', RetoAnalisis.unidades_relacionar, ids),
        _panel_m2m('Desagregación del análisis (1.6)', RetoAnalisis.desagregaciones, ids),
        _panel_m2m('Tratamiento analítico (1.10)', RetoAnalisis.tratamientos, ids),
        _panel_m2m('Resultado mínimo viable (1.11)', RetoAnalisis.resultados_mvp, ids),
        _panel_m2m('Capacidades del SIPDI (1.12)', RetoAnalisis.capacidades_sipdi, ids),
        _panel_texto('Variables a enriquecer (1.13)', retos, 'variables_enriquecer'),
        _panel_m2m('Enfoques diferenciales (1.15)', RetoAnalisis.enfoques_diferenciales, ids),
    ]
    unidad_principal = _filas(list(
        retos.exclude(unidad_principal__isnull=True)
        .values_list('unidad_principal__etiqueta')
        .annotate(total=Count('id')).order_by('-total')))
    paneles_f2.insert(0, {'titulo': 'Unidad principal de análisis (1.5)', 'filas': unidad_principal})

    paneles_f3_acceso = [
        _panel_choice('¿Contiene datos personales?', seguridades, 'contiene_datos_personales'),
        _panel_m2m('Variables que identifican personas', EvaluacionSeguridad.variables_identificadoras, ids),
        _panel_m2m('Información de especial cuidado', EvaluacionSeguridad.datos_sensibles, ids),
        _panel_choice('¿Base legítima para acceder y utilizar?', seguridades, 'base_legitima'),
    ]
    paneles_f3_uso = [
        _panel_choice('Debida diligencia del proveedor', seguridades, 'proveedor_base_legitima'),
        _panel_choice('Finalidad específica y restringida', seguridades, 'finalidad_especifica'),
        _panel_choice('Datos necesarios y no excesivos', seguridades, 'datos_necesarios'),
        _panel_choice('Conservación limitada en el tiempo', seguridades, 'conservacion_limitada'),
        _panel_choice('Salvaguardas técnicas/administrativas/físicas', seguridades, 'salvaguardas'),
        _panel_choice('Transparencia con titulares y público', seguridades, 'transparencia'),
        _panel_choice('Riesgos de publicación desproporcionados', seguridades, 'riesgos_publicacion'),
        _panel_choice('Terceros con estándares equivalentes', seguridades, 'terceros_estandares'),
        _panel_choice('Riesgo de daño a personas o grupos', seguridades, 'riesgo_dano'),
    ]
    paneles_f3_medidas = [
        _panel_m2m('Medidas de protección aplicadas', EvaluacionSeguridad.medidas_proteccion, ids),
        _panel_choice('¿Susceptible a fraude o corrupción?', seguridades, 'susceptible_fraude'),
    ]
    paneles_semaforo = [
        _panel_choice('Confidencialidad', seguridades, 'confidencialidad', semaforo=True),
        _panel_choice('Integridad', seguridades, 'integridad', semaforo=True),
        _panel_choice('Disponibilidad', seguridades, 'disponibilidad', semaforo=True),
        _panel_choice('Criticidad (calculada)', seguridades, 'criticidad', semaforo=True),
    ]
    paneles_clasificacion = [
        _panel_choice('Aplica: Pública', seguridades, 'aplica_publica'),
        _panel_choice('Aplica: Semiprivada', seguridades, 'aplica_semiprivada'),
        _panel_choice('Aplica: Privada / restringida', seguridades, 'aplica_privada'),
        _panel_choice('Aplica: Sensible', seguridades, 'aplica_sensible'),
        _panel_choice('Aplica: Reservada', seguridades, 'aplica_reservada'),
    ]

    context = {
        'alcance': alcance,
        'total_inscritos': todas.count(),
        'total_completadas': completadas.count(),
        'total_en_progreso': todas.count() - completadas.count(),
        'total_organismos': todas.values('organismo').distinct().count(),
        'total_fuentes': fuentes.exclude(nombre='').count(),
        'total_retos': retos.exclude(nombre='').count(),
        'total_variables': variables.count(),
        'total_identificadores': variables.filter(es_identificador=True).count(),
        'criticos': seguridades.filter(criticidad='rojo').count(),
        'con_datos_personales': seguridades.filter(contiene_datos_personales='si').count(),
        'embudo': embudo,
        'fases': wz.FASES,
        'organismos': organismos,
        'fuentes_top': fuentes_top,
        'requeridas_top': requeridas_top,
        'paneles_f1': [p for p in paneles_f1 if p['filas']],
        'paneles_f2': [p for p in paneles_f2 if p['filas']],
        'paneles_f3_acceso': [p for p in paneles_f3_acceso if p['filas']],
        'paneles_f3_uso': [p for p in paneles_f3_uso if p['filas']],
        'paneles_f3_medidas': [p for p in paneles_f3_medidas if p['filas']],
        'paneles_semaforo': [p for p in paneles_semaforo if p['filas']],
        'paneles_clasificacion': [p for p in paneles_clasificacion if p['filas']],
        'ultimas': _con_criticidad(inscripciones),
        'fuentes_detalle': _fuentes_detalle(fuentes, ids),
    }
    return render(request, 'taller/dashboard.html', context)


@taller_staff_required
def presentacion(request):
    """Presentación en diapositivas del análisis integral F1-F2-F3 — cifras en vivo desde la BD."""
    from collections import Counter

    inscripciones = Inscripcion.objects.count()
    fuentes = FuenteDatos.objects.all()
    n_fuentes = fuentes.count()
    retos = RetoAnalisis.objects.count()
    evals = EvaluacionSeguridad.objects.all()
    n_evals = evals.count()
    organismos = Inscripcion.objects.values('organismo').distinct().count()
    org_con_fuentes = Inscripcion.objects.filter(fuentes__isnull=False).values('organismo').distinct().count()
    fuentes_sin_f3 = fuentes.filter(seguridad__isnull=True).count()

    # F1 · top fuentes por volumen
    top_vol_qs = (fuentes.exclude(num_registros__isnull=True)
                  .order_by('-num_registros').values_list('nombre', 'num_registros')[:9])
    max_vol = top_vol_qs[0][1] if top_vol_qs else 1
    top_vol = [{'nombre': n, 'valor': v, 'pct': max(round(v / max_vol * 100, 1), 0.4)}
               for n, v in top_vol_qs]

    # F1 · escalera de madurez
    n_dicc = fuentes.filter(tiene_diccionario='si').count()
    n_val = fuentes.filter(tiene_validacion='si').count()

    def dist(field):
        d = {'si': 0, 'no': 0, 'no_se': 0, 'no_aplica': 0, '': 0}
        for v, c in evals.values_list(field).annotate(c=Count('id')):
            d[v] = c
        return d

    def habeas_row(label, field, invertir_no=False):
        """Fila de la matriz habeas data. invertir_no: cuando 'No' es la respuesta favorable (riesgos_publicacion)."""
        d = dist(field)
        total = d['si'] + d['no'] + d['no_se'] + d['no_aplica'] or 1
        if invertir_no:
            fav, desf = d['no'], d['si']
        else:
            fav, desf = d['si'], d['no']
        return {'label': label, 'fav': fav, 'ns': d['no_se'], 'desf': desf, 'na': d['no_aplica'],
                'p_fav': round(fav / total * 100), 'p_ns': round(d['no_se'] / total * 100),
                'p_desf': round(desf / total * 100), 'p_na': round(d['no_aplica'] / total * 100)}

    habeas = [
        habeas_row('Base legítima', 'base_legitima'),
        habeas_row('Finalidad específica', 'finalidad_especifica'),
        habeas_row('Datos necesarios (minimización)', 'datos_necesarios'),
        habeas_row('Conservación limitada', 'conservacion_limitada'),
        habeas_row('Salvaguardas', 'salvaguardas'),
        habeas_row('Transparencia', 'transparencia'),
        habeas_row('Riesgos de publicación', 'riesgos_publicacion', invertir_no=True),
        habeas_row('Terceros con estándares', 'terceros_estandares'),
    ]

    # F3 · CIA
    def semaforo_row(label, field):
        d = dist(field)
        total = d['verde'] + d['amarillo'] + d['rojo'] or 1
        return {'label': label, 'v': d['verde'], 'a': d['amarillo'], 'r': d['rojo'],
                'p_v': round(d['verde'] / total * 100), 'p_a': round(d['amarillo'] / total * 100),
                'p_r': round(d['rojo'] / total * 100)}

    cia = [semaforo_row('Confidencialidad', 'confidencialidad'),
           semaforo_row('Integridad', 'integridad'),
           semaforo_row('Disponibilidad', 'disponibilidad'),
           semaforo_row('Criticidad', 'criticidad')]

    fuentes_rojas = (evals.filter(criticidad='rojo').select_related('fuente')
                     .values('fuente__nombre', 'confidencialidad', 'integridad', 'disponibilidad'))

    # F3 · exposición
    n_dp = evals.filter(contiene_datos_personales='si').count()
    n_sens = evals.filter(datos_sensibles__isnull=False).distinct().count()
    n_fraude = evals.filter(susceptible_fraude='si').count()

    def m2m_bars(rel, top=8):
        c = Counter()
        for e in evals.prefetch_related(rel):
            for o in getattr(e, rel).all():
                c[o.etiqueta] += 1
        items = c.most_common(top)
        mx = items[0][1] if items else 1
        return [{'label': k, 'valor': v, 'pct': max(round(v / mx * 100), 4)} for k, v in items]

    identificadoras = m2m_bars('variables_identificadoras')
    medidas = m2m_bars('medidas_proteccion', top=10)
    sensibles = Counter()
    for e in evals.prefetch_related('datos_sensibles'):
        for o in e.datos_sensibles.all():
            sensibles[o.etiqueta] += 1

    # F3 · clasificación
    clasif = {
        'publica': evals.filter(aplica_publica='si').count(),
        'semiprivada': evals.filter(aplica_semiprivada='si').count(),
        'privada': evals.filter(aplica_privada='si').count(),
        'sensible': evals.filter(aplica_sensible='si').count(),
        'reservada': evals.filter(aplica_reservada='si').count(),
    }

    context = {
        'inscripciones': inscripciones,
        'n_fuentes': n_fuentes,
        'retos': retos,
        'n_evals': n_evals,
        'organismos': organismos,
        'org_con_fuentes': org_con_fuentes,
        'fuentes_sin_f3': fuentes_sin_f3,
        'top_vol': top_vol,
        'n_dicc': n_dicc,
        'n_val': n_val,
        'habeas': habeas,
        'cia': cia,
        'fuentes_rojas': fuentes_rojas,
        'n_dp': n_dp,
        'n_sens': n_sens,
        'n_fraude': n_fraude,
        'identificadoras': identificadoras,
        'medidas': medidas,
        'sensibles': sensibles.most_common(),
        'clasif': clasif,
    }

    # ---- Datos F2 vivos: taxonomía por línea, franjas, brechas ----
    casos = _clasificar_retos()
    clasificables = [c for c in casos if c['primaria']]
    context['n_casos_f2'] = len(clasificables)
    context['n_multilinea'] = sum(1 for c in clasificables if c['n_lineas'] > 1)
    context['n_cuatro'] = sum(1 for c in clasificables if c['n_lineas'] == 4)

    # taxonomía: conteo + 3 ejemplos por línea
    taxonomia = []
    for code in ['CPS', 'ADA', 'FM', 'RR']:
        grupo = [c for c in clasificables if c['primaria'] == code]
        taxonomia.append({
            'code': code, 'label': LINEAS_F2[code], 'count': len(grupo),
            'ejemplos': [{'id': c['id'], 'nombre': c['nombre']} for c in grupo[:3]],
        })
    context['taxonomia'] = taxonomia

    # franjas de secuenciación
    franjas = {}
    for f in [1, 2, 3]:
        grupo = [c for c in clasificables if c['franja'] == f]
        franjas[f] = {'n': len(grupo),
                      'ids': ' · '.join(f"R{c['id']}" for c in grupo)}
    context['franjas'] = franjas

    # brechas F1: EDE y RUD desde la BD
    ede = fuentes.filter(nombre__icontains='Especializadas').first() or \
        fuentes.filter(nombre__icontains='EDE').first()
    rud = fuentes.filter(nombre__icontains='Damnificados').first() or \
        fuentes.filter(nombre__icontains='RUD').first()
    context['ede_reg'] = ede.num_registros if ede else None
    context['rud_reg'] = rud.num_registros if rud else None
    context['ede_bloquea'] = ' · '.join(
        f"R{c['id']}" for c in casos
        if any('especializ' in fr['nombre'].lower() or 'ede' in fr['nombre'].lower()
               for fr in c['reqs']))
    context['rud_casos'] = sum(
        1 for c in casos
        if any('rud' in fr['nombre'].lower() or 'damnific' in fr['nombre'].lower()
               for fr in c['reqs']))

    # capa Gold (constantes de arquitectura SIPDI)
    context['n_gold'] = 9
    context['n_gold_cat'] = 2
    context['n_gold_int'] = 0

    # CIA y protección: conteos vivos
    context['n_rojas'] = len(fuentes_rojas)
    context['cia_int_rojos'] = next((c['r'] for c in cia if c['label'] == 'Integridad'), 0)
    medidas_map = {m['label']: m['valor'] for m in medidas}
    context['n_cifrado'] = medidas_map.get('Cifrado de datos', 0)
    context['n_anonim'] = (medidas_map.get('Eliminación de identificadores', 0) +
                           medidas_map.get('Sustitución de identificadores por códigos', 0))
    context['n_logs'] = medidas_map.get('Registros de auditoría (logs)', 0)

    return render(request, 'taller/presentacion.html', context)


@taller_staff_required
def nube_svg(request, nombre):
    """Sirve las nubes de palabras SVG desde el paquete (sin depender de collectstatic)."""
    from pathlib import Path
    if nombre not in {'global', 'cps', 'ada', 'fm', 'rr'}:
        raise Http404
    ruta = Path(__file__).resolve().parent / 'nubes' / f'nube_{nombre}.svg'
    if not ruta.exists():
        raise Http404
    return FileResponse(open(ruta, 'rb'), content_type='image/svg+xml')


LINEAS_F2 = {'CPS': 'Caracterización poblacional y socioeconómica',
             'ADA': 'Afectaciones, daños y activos',
             'FM': 'Focalización microterritorial',
             'RR': 'Respuesta y recuperación'}

_CODEBOOK_F2 = {
    'CPS': ['persona', 'hogar', 'poblacion', 'sisben', 'grupo poblacional',
            'caracterizacion socioeconomica', 'damnificado', 'afectado',
            'deportista', 'colaborador', 'quienes', 'cuantos', 'edad', 'sexo',
            'etnia', 'discapacidad', 'victima', 'familia', 'comunidad',
            'estudiante', 'menor', 'adulto mayor', 'demograf', 'socioeconom'],
    'ADA': ['daño', 'afectacion', 'edificacion', 'predio', 'infraestructura',
            'inmueble', 'equipamiento', 'avaluo', 'valor de la afectacion',
            'rcd', 'residuo', 'terminal', 'estacion', 'nivel de daño', 'colapso',
            'estructura', 'vivienda', 'alojamiento', 'dano', 'escombro',
            'demolicion', 'sismo', 'terremoto'],
    'FM': ['donde', 'localiz', 'prioriz', 'territorio', 'albergue', 'cobertura',
           'brecha', 'georreferenc', 'tramo', 'microterritorio', 'comuna',
           'barrio', 'mapa', 'ubicacion', 'zona', 'sector', 'espacio publico',
           'corredor', 'focaliz', 'concentracion', 'densidad'],
    'RR': ['ayuda', 'apoyo', 'donacion', 'recuperacion', 'reencuentro',
           'respuesta', 'asistencia', 'seguimiento', 'oferta', 'albergue',
           'alojamiento temporal', 'brote', 'salud', 'atencion', 'subsidio',
           'compensac', 'rehabilitac', 'reconstruc', 'animal', 'empleo',
           'empresa', 'negocio'],
}


def _clasificar_retos():
    """Clasificación viva de los retos F2: línea primaria/secundarias por codebook,
    fuentes F1 propias, fuentes requeridas con llave y limitación de integración."""
    retos = (RetoAnalisis.objects
             .select_related('inscripcion__organismo', 'unidad_principal')
             .prefetch_related('fuentes_requeridas', 'inscripcion__fuentes',
                               'unidades_relacionar'))
    casos = []
    for r in retos:
        ins = r.inscripcion
        texto = ' '.join(filter(None, [
            r.nombre, r.analisis_propuesto, r.resultado_esperado,
            r.decision_apoya,
            r.unidad_principal.etiqueta if r.unidad_principal else '',
            ' '.join(o.etiqueta for o in r.unidades_relacionar.all()),
            r.variables_enriquecer])).lower()
        scores = {l: sum(texto.count(k) for k in kws)
                  for l, kws in _CODEBOOK_F2.items()}
        orden = sorted(scores.items(), key=lambda x: -x[1])
        primaria = orden[0][0] if orden[0][1] > 0 else ''
        secundarias = [l for l, s in orden[1:] if s > 0]

        reqs = list(r.fuentes_requeridas.all())
        sin_llave = [fr.nombre for fr in reqs if not fr.llave_relacion.strip()]
        if primaria == '':
            limitacion = 'Reto sin contenido: no evaluable.'
        elif not reqs:
            limitacion = 'No declaró fuentes requeridas: el insumo no está especificado.'
        elif sin_llave:
            limitacion = ('Sin llave de cruce en: ' + '; '.join(sin_llave) +
                          '. Requiere geocodificación, relación espacial o conciliación.')
        else:
            limitacion = 'Con llave declarada en todas las fuentes requeridas.'

        # Franja de secuenciación por disponibilidad del dato
        def _tipo_llave(l):
            l = (l or '').lower()
            if any(k in l for k in ['documento', 'cedula', 'cédula', 'nit', 'num_doc', 'id']):
                return 'identidad'
            if any(k in l for k in ['predio', 'direccion', 'dirección', 'coordenada',
                                    'latitud', 'longitud', 'barrio', 'comuna',
                                    'matricula', 'matrícula', 'npn', 'ficha']):
                return 'espacial'
            return 'declarativa'
        if primaria == '':
            franja = 0
        elif not reqs and not ins.fuentes.exists():
            franja = 3
        elif not reqs:
            franja = 2
        else:
            tipos = [_tipo_llave(fr.llave_relacion) for fr in reqs]
            if all(t == 'identidad' for t in tipos) and r.variables_enriquecer.strip():
                franja = 1
            elif any(t == 'declarativa' for t in tipos):
                franja = 3
            else:
                franja = 2

        casos.append({
            'id': r.id,
            'organismo': ins.organismo.nombre if ins.organismo else '',
            'nombre': r.nombre,
            'analisis': r.analisis_propuesto,
            'decision': r.decision_apoya,
            'primaria': primaria,
            'primaria_label': LINEAS_F2.get(primaria, 'Sin clasificar'),
            'secundarias': secundarias,
            'n_lineas': 1 + len(secundarias) if primaria else 0,
            'fuentes_f1': [f.nombre for f in ins.fuentes.all()],
            'reqs': [{'nombre': fr.nombre, 'org': fr.organismo,
                      'vars': fr.variables, 'llave': fr.llave_relacion}
                     for fr in reqs],
            'enriquecer': r.variables_enriquecer,
            'limitacion': limitacion,
            'franja': franja,
        })
    return casos


@taller_staff_required
def casos_f2(request):
    """Tabla viva de los casos F2 clasificados por línea — consultable desde el dashboard."""
    from collections import Counter
    casos = _clasificar_retos()
    por_linea = Counter(c['primaria'] or 'SIN' for c in casos)
    context = {
        'casos': sorted(casos, key=lambda c: (c['primaria'] or 'Z', c['organismo'])),
        'n_casos': len(casos),
        'por_linea': por_linea,
        'lineas': LINEAS_F2,
    }
    return render(request, 'taller/casos_f2.html', context)


@taller_staff_required
def casos_f2_excel(request):
    """Exporta la tabla de casos F2 clasificados a Excel."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    casos = _clasificar_retos()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Casos F2'
    header = ['ID', 'Organismo', 'Reto de análisis', 'Línea primaria',
              'Líneas secundarias', 'Nº líneas', 'Fuentes F1 propias',
              'Fuentes requeridas', 'Llave de cruce', 'Variables a enriquecer',
              'Limitación para el análisis']
    ws.append(header)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF', size=10)
        celda.fill = PatternFill('solid', fgColor='004B87')
    for c in sorted(casos, key=lambda x: (x['primaria'] or 'Z', x['organismo'])):
        ws.append([
            f"R{c['id']}", c['organismo'], c['nombre'],
            c['primaria_label'], '+'.join(c['secundarias']) or '—',
            c['n_lineas'] or '—',
            '\n'.join('• ' + f for f in c['fuentes_f1']) or '—',
            '\n'.join(f"• {fr['nombre']} ({fr['org']})" for fr in c['reqs']) or '—',
            '\n'.join(f"• {fr['nombre']}: {fr['llave'] or 'SIN LLAVE'}"
                      for fr in c['reqs']) or '—',
            c['enriquecer'] or '—', c['limitacion']])
    for idx, col in enumerate(ws.columns, 1):
        ancho = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(ancho + 2, 50)
    ws.freeze_panes = 'A2'

    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = 'attachment; filename="casos_f2_clasificados.xlsx"'
    wb.save(respuesta)
    return respuesta


@taller_staff_required
def reporte_excel(request):
    """Informe claro multi-hoja: separa ficha/inscripción, fuentes F1, retos F2 y
    evaluaciones F3 para que no se confundan las unidades de conteo."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    def estilo_header(ws):
        for celda in ws[1]:
            celda.font = Font(bold=True, color='FFFFFF', size=10)
            celda.fill = PatternFill('solid', fgColor='004B87')
            celda.alignment = Alignment(vertical='center', wrap_text=True)
        ws.freeze_panes = 'A2'

    def anchos(ws, tope=48):
        for idx, col in enumerate(ws.columns, 1):
            ancho = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(ancho + 2, tope)

    def m2m(obj, campo):
        return '; '.join(o.etiqueta for o in getattr(obj, campo).all()) if obj else ''

    def disp(obj, campo):
        return getattr(obj, f'get_{campo}_display')() if obj and getattr(obj, campo) else ''

    inscripciones = (Inscripcion.objects
                     .select_related('organismo', 'suborganismo', 'user',
                                     'reto', 'reto__unidad_principal')
                     .prefetch_related(
                         'fuentes', 'fuentes__seguridad',
                         'fuentes__relaciones', 'fuentes__instrumentos',
                         'fuentes__mecanismos', 'fuentes__periodicidades',
                         'fuentes__desagregaciones', 'fuentes__unidades_analisis',
                         'fuentes__niveles_georref', 'fuentes__problemas_calidad',
                         'fuentes__fuentes_origen', 'fuentes__variables',
                         'fuentes__seguridad__variables_identificadoras',
                         'fuentes__seguridad__datos_sensibles',
                         'fuentes__seguridad__medidas_proteccion',
                         'reto__unidades_relacionar', 'reto__desagregaciones',
                         'reto__tratamientos', 'reto__resultados_mvp',
                         'reto__capacidades_sipdi', 'reto__enfoques_diferenciales',
                         'reto__fuentes_requeridas')
                     .order_by('organismo__nombre', 'id'))
    casos = {c['id']: c for c in _clasificar_retos()}

    wb = Workbook()

    # ---- Hoja 1 · Resumen / glosario ----
    ws = wb.active
    ws.title = 'Resumen'
    n_fuentes = FuenteDatos.objects.count()
    n_retos = RetoAnalisis.objects.count()
    n_evals = EvaluacionSeguridad.objects.count()
    fichas_con_f1 = inscripciones.filter(fuentes__isnull=False).distinct().count()
    ws.append(['Concepto', 'Qué es', 'Cantidad'])
    for fila in [
        ['Ficha / Inscripción', 'El formulario diligenciado por un organismo (una sola cosa: la inscripción ES la ficha). Contiene identificación + F1 + F2 + F3.', inscripciones.count()],
        ['Fuente F1', 'Cada conjunto de datos caracterizado. Una ficha puede tener 0, 1 o varias fuentes.', n_fuentes],
        ['Reto F2', 'El reto de análisis propuesto. Máximo uno por ficha.', n_retos],
        ['Evaluación F3', 'La evaluación de seguridad de una fuente. Máximo una por fuente F1.', n_evals],
        ['', '', ''],
        ['Fichas CON al menos una fuente F1', f'Por eso {n_fuentes} fuentes no son {n_fuentes} fichas.', fichas_con_f1],
        ['Fichas SIN fuente F1', 'Diligenciaron solo F2 u otra sección.', inscripciones.count() - fichas_con_f1],
        ['Fuentes SIN evaluación F3', 'Fuentes caracterizadas sin evaluar seguridad.', n_fuentes - n_evals],
    ]:
        ws.append(fila)
    estilo_header(ws)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
    anchos(ws, 70)

    # ---- Hoja 2 · Fichas (una fila por inscripción — TODOS los campos) ----
    ws = wb.create_sheet('Fichas')
    ws.append(['ID ficha', 'UUID', 'Usuario', 'Estado', 'Inició por',
               'Último paso', 'Creado', 'Actualizado',
               'Organismo', 'Suborganismo', 'Equipo responsable',
               'Nombre completo', 'Correo', 'Teléfono',
               'Nº fuentes F1', 'Tiene reto F2', 'Nº fuentes con F3'])
    for i in inscripciones:
        fuentes_i = list(i.fuentes.all())
        ws.append([
            i.id, str(i.uuid), i.user.username if i.user else '',
            i.get_estado_display(), i.get_ficha_inicial_display(),
            i.paso_actual,
            i.creado.strftime('%Y-%m-%d %H:%M') if i.creado else '',
            i.actualizado.strftime('%Y-%m-%d %H:%M') if i.actualizado else '',
            i.organismo.nombre if i.organismo else '',
            i.suborganismo.nombre if i.suborganismo else '',
            i.equipo_responsable, i.nombre_completo, i.correo, i.telefono,
            len(fuentes_i),
            'Sí' if getattr(i, 'reto', None) else 'No',
            sum(1 for f in fuentes_i if getattr(f, 'seguridad', None)),
        ])
    estilo_header(ws); anchos(ws)

    # ---- Hoja 3 · Fuentes F1 (una fila por fuente — TODOS los campos) ----
    ws = wb.create_sheet('Fuentes F1')
    ws.append(['ID fuente', 'ID ficha', 'Organismo',
               '1.1 Nombre', '1.2 Descripción',
               '1.3 Relaciones', '1.3 Relación (otra)',
               '1.4 Instrumentos', '1.4 Instrumento (otro)',
               '1.4 Mecanismos', '1.4 Mecanismo (otro)',
               '1.5 Fecha inicio', '1.6 Última actualización',
               '1.7 Periodicidades', '1.7 Periodicidad (otra)',
               '1.8 Desagregaciones', '1.8 Desagregación (otra)',
               '1.9 Nº registros', '1.9 Registros (fecha)', '1.9 Registros desconocido',
               '1.10 Tiene diccionario', '1.10 Diccionario (ubicación)', '1.10 Diccionario (archivo)',
               '1.11 Rol', '1.11 Fuentes de origen',
               '1.12 Decisiones que apoya',
               '1.13 Unidades de análisis', '1.13 Unidad (otra)',
               '1.14 Georreferenciación', '1.15 Niveles georref', '1.15 Nivel georref (otro)',
               '1.16 Problemas de calidad', '1.16 Problema (otro)',
               '1.17 Tiene validación', '1.17 Validación (descripción)',
               'Variables (nombre:tipo:desc:identificador)'])
    for i in inscripciones:
        for f in i.fuentes.all():
            f_orig = '; '.join(f'{o.nombre} ({o.organismo_productor})'
                               for o in f.fuentes_origen.all())
            f_vars = '; '.join(
                f'{v.nombre}:{v.get_tipo_display()}:{v.descripcion or ""}'
                f'{" [ID]" if v.es_identificador else ""}'
                for v in f.variables.all())
            ws.append([
                f.id, i.id, i.organismo.nombre if i.organismo else '',
                f.nombre, f.descripcion,
                m2m(f, 'relaciones'), f.relacion_otra,
                m2m(f, 'instrumentos'), f.instrumento_otro,
                m2m(f, 'mecanismos'), f.mecanismo_otro,
                f.fecha_inicio, f.fecha_actualizacion,
                m2m(f, 'periodicidades'), f.periodicidad_otra,
                m2m(f, 'desagregaciones'), f.desagregacion_otra,
                f.num_registros, f.num_registros_fecha,
                'Sí' if f.num_registros_desconocido else '',
                disp(f, 'tiene_diccionario'), f.diccionario_ubicacion,
                f.archivo_diccionario.name if f.archivo_diccionario else '',
                disp(f, 'rol'), f_orig,
                f.decisiones_apoya,
                m2m(f, 'unidades_analisis'), f.unidad_analisis_otra,
                disp(f, 'georreferenciacion'), m2m(f, 'niveles_georref'), f.nivel_georref_otro,
                m2m(f, 'problemas_calidad'), f.problema_calidad_otro,
                disp(f, 'tiene_validacion'), f.validacion_descripcion,
                f_vars,
            ])
    estilo_header(ws); anchos(ws)

    # ---- Hoja 4 · Retos F2 (una fila por reto — TODOS los campos) ----
    ws = wb.create_sheet('Retos F2')
    ws.append(['ID reto', 'ID ficha', 'Organismo',
               '1.1 Nombre del reto', '1.2 Análisis propuesto',
               '1.3 Resultado esperado', '1.4 Decisión que apoya',
               '1.5 Unidad principal', '1.5 Unidad (otra)',
               '1.6 Unidades a relacionar', '1.6 Unidad (otra)', '1.6 Vínculo entre unidades',
               '1.6 Desagregaciones', '1.6 Desagregación (otra)',
               '1.10 Tratamientos', '1.10 Tratamiento (otro)',
               '1.11 Resultados MVP', '1.11 Resultado (otro)',
               '1.12 Capacidades SIPDI', '1.12 Capacidad (otra)',
               '1.13 Variables a enriquecer', '1.13 Variables (otro)',
               '1.14 Quién usará el resultado',
               '1.15 Enfoques diferenciales', '1.15 Enfoque (otro)',
               'Req1 fuente', 'Req1 custodio', 'Req1 variables (1.8)', 'Req1 llave (1.9)',
               'Req2 fuente', 'Req2 custodio', 'Req2 variables (1.8)', 'Req2 llave (1.9)',
               'Req3 fuente', 'Req3 custodio', 'Req3 variables (1.8)', 'Req3 llave (1.9)',
               'Línea primaria', 'Líneas secundarias', 'Franja', 'Limitación'])
    for i in inscripciones:
        r = getattr(i, 'reto', None)
        if not r:
            continue
        c = casos.get(r.id, {})
        reqs = list(r.fuentes_requeridas.all())[:3]
        req_cols = [campo for q in reqs
                    for campo in (q.nombre, q.organismo, q.variables, q.llave_relacion)]
        req_cols += [''] * (12 - len(req_cols))
        ws.append([
            r.id, i.id, i.organismo.nombre if i.organismo else '',
            r.nombre, r.analisis_propuesto, r.resultado_esperado, r.decision_apoya,
            r.unidad_principal.etiqueta if r.unidad_principal else '', r.unidad_principal_otra,
            m2m(r, 'unidades_relacionar'), r.unidades_relacionar_otra, r.vinculo_unidades,
            m2m(r, 'desagregaciones'), r.desagregacion_otra,
            m2m(r, 'tratamientos'), r.tratamiento_otro,
            m2m(r, 'resultados_mvp'), r.resultado_mvp_otro,
            m2m(r, 'capacidades_sipdi'), r.capacidad_sipdi_otra,
            r.variables_enriquecer, r.variables_enriquecer_otro,
            r.usuario_resultado,
            m2m(r, 'enfoques_diferenciales'), r.enfoque_diferencial_otro,
            *req_cols,
            c.get('primaria_label', ''), '+'.join(c.get('secundarias', [])) or '—',
            {0: 'Sin evaluar', 1: 'Franja 1', 2: 'Franja 2', 3: 'Franja 3'}.get(c.get('franja'), '—'),
            c.get('limitacion', ''),
        ])
    estilo_header(ws); anchos(ws)

    # ---- Hoja 5 · Evaluaciones F3 (una fila por evaluación — TODOS los campos) ----
    ws = wb.create_sheet('Evaluaciones F3')
    ws.append(['ID fuente', 'ID ficha', 'Organismo', 'Fuente',
               'Datos personales',
               'Variables identificadoras', 'Variable identificadora (otra)',
               'Datos sensibles', 'Dato sensible (otro)',
               'Base legítima', 'Base legítima (por qué)',
               'Proveedor base legítima', 'Proveedor base legítima (por qué)',
               'Finalidad específica', 'Finalidad específica (por qué)',
               'Datos necesarios', 'Datos necesarios (por qué)',
               'Conservación limitada', 'Conservación limitada (por qué)',
               'Salvaguardas', 'Salvaguardas (por qué)',
               'Transparencia', 'Transparencia (por qué)',
               'Riesgos publicación', 'Riesgos publicación (por qué)',
               'Terceros estándares', 'Terceros estándares (por qué)',
               'Riesgo de daño', 'Riesgo de daño (por qué)',
               'Medidas de protección', 'Medida protección (otra)',
               'Susceptible a fraude', 'Organismo custodio',
               'Confidencialidad', 'Integridad', 'Disponibilidad', 'Criticidad',
               'Aplica pública', 'Aplica semiprivada', 'Aplica privada',
               'Aplica sensible', 'Aplica reservada'])
    for i in inscripciones:
        for f in i.fuentes.all():
            s = getattr(f, 'seguridad', None)
            if not s:
                continue
            ws.append([
                f.id, i.id, i.organismo.nombre if i.organismo else '', f.nombre,
                disp(s, 'contiene_datos_personales'),
                m2m(s, 'variables_identificadoras'), s.variable_identificadora_otra,
                m2m(s, 'datos_sensibles'), s.dato_sensible_otro,
                disp(s, 'base_legitima'), s.base_legitima_porque,
                disp(s, 'proveedor_base_legitima'), s.proveedor_base_legitima_porque,
                disp(s, 'finalidad_especifica'), s.finalidad_especifica_porque,
                disp(s, 'datos_necesarios'), s.datos_necesarios_porque,
                disp(s, 'conservacion_limitada'), s.conservacion_limitada_porque,
                disp(s, 'salvaguardas'), s.salvaguardas_porque,
                disp(s, 'transparencia'), s.transparencia_porque,
                disp(s, 'riesgos_publicacion'), s.riesgos_publicacion_porque,
                disp(s, 'terceros_estandares'), s.terceros_estandares_porque,
                disp(s, 'riesgo_dano'), s.riesgo_dano_porque,
                m2m(s, 'medidas_proteccion'), s.medida_proteccion_otra,
                disp(s, 'susceptible_fraude'), s.organismo_custodio,
                disp(s, 'confidencialidad'), disp(s, 'integridad'),
                disp(s, 'disponibilidad'), disp(s, 'criticidad'),
                disp(s, 'aplica_publica'), disp(s, 'aplica_semiprivada'),
                disp(s, 'aplica_privada'), disp(s, 'aplica_sensible'),
                disp(s, 'aplica_reservada'),
            ])
    estilo_header(ws); anchos(ws)

    # ---- Hoja 6 · Variables llave candidatas al SIPDI ----
    ws = wb.create_sheet('Variables llave SIPDI')
    ws.append(['Variable', 'Rol', 'Nº retos', 'Retos que la piden',
               'Nº fuentes donde es llave', 'Fuentes que la usan como llave'])
    for v in _variables_llave():
        ws.append([v['variable'], v['rol'], v['n_retos'],
                   ' · '.join(f'R{r}' for r in v['retos']),
                   v['n_llaves'], ' · '.join(v['fuentes_llave'])])
    estilo_header(ws); anchos(ws, 60)

    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = 'attachment; filename="informe_taller.xlsx"'
    wb.save(respuesta)
    return respuesta


def _variables_llave():
    """Consolida las variables llave candidatas al SIPDI.

    Cruza dos fuentes del instrumento:
    - F2 1.13 `variables_enriquecer`: variables que cada reto quiere enriquecer/
      validar mediante el SIPDI (demanda).
    - F2 1.9 `llave_relacion` de cada fuente requerida: la llave de cruce
      declarada (oferta de interoperabilidad).
    Retorna filas ordenadas por demanda con el rol de cada variable.
    """
    import re
    from collections import defaultdict

    def split_vars(txt):
        partes = re.split(r'[,;]|\s+y\s+', (txt or '').lower())
        return [p.strip().strip('.') for p in partes if p.strip()]

    def canon_llave(t):
        t = (t or '').lower()
        if re.search(r'documento|c[eé]dula|num_doc|nit', t):
            return 'documento/NIT'
        if re.search(r'predio|predial|npn|matr[ií]cula', t):
            return 'predio'
        if re.search(r'vivienda|hogar|ficha', t):
            return 'vivienda/hogar'
        if re.search(r'coordenada|latitud|longitud|geo|placeid', t):
            return 'coordenadas/geo'
        if re.search(r'direcci[oó]n', t):
            return 'dirección'
        if re.search(r'comuna|barrio', t):
            return 'comuna/barrio'
        return t.strip()[:30]

    # Demanda: variable → retos que la piden (1.13)
    retos_por_var = defaultdict(set)
    for r in RetoAnalisis.objects.exclude(variables_enriquecer=''):
        for v in split_vars(r.variables_enriquecer):
            retos_por_var[v].add(r.id)

    # Oferta: llave canónica → fuentes que la usan (1.9)
    fuentes_por_llave = defaultdict(set)
    retos_por_llave = defaultdict(set)
    for q in FuenteRequerida.objects.exclude(llave_relacion=''):
        c = canon_llave(q.llave_relacion)
        fuentes_por_llave[c].add(q.nombre)
        retos_por_llave[c].add(q.reto_id)

    # Mapear variables de enriquecimiento a su llave canónica equivalente
    def llave_de(var):
        return canon_llave(var) if canon_llave(var) in fuentes_por_llave else None

    filas = []
    for var, retos in retos_por_var.items():
        lk = llave_de(var)
        filas.append({
            'variable': var,
            'rol': 'Llave de cruce + enriquecimiento' if lk else 'Enriquecimiento',
            'n_retos': len(retos),
            'retos': sorted(retos),
            'n_llaves': len(fuentes_por_llave.get(lk, ())) if lk else 0,
            'fuentes_llave': sorted(fuentes_por_llave.get(lk, ())) if lk else [],
        })
    # Llaves de cruce que nadie pidió como enriquecimiento (solo oferta)
    for lk, fuentes in fuentes_por_llave.items():
        if not any(llave_de(v) == lk for v in retos_por_var):
            filas.append({
                'variable': lk,
                'rol': 'Solo llave de cruce',
                'n_retos': len(retos_por_llave[lk]),
                'retos': sorted(retos_por_llave[lk]),
                'n_llaves': len(fuentes),
                'fuentes_llave': sorted(fuentes),
            })
    filas.sort(key=lambda x: (-x['n_retos'], -x['n_llaves'], x['variable']))
    return filas


def _con_criticidad(inscripciones):
    """Adjunta a cada inscripción la peor criticidad entre sus fuentes (F3 por fuente)."""
    orden = {'rojo': 3, 'amarillo': 2, 'verde': 1, '': 0}
    for i in inscripciones:
        peor = ''
        for f in i.fuentes.all():
            seg = getattr(f, 'seguridad', None)
            if seg and orden.get(seg.criticidad, 0) > orden.get(peor, 0):
                peor = seg.criticidad
        i.crit_peor = peor
    return inscripciones


def _export_filas(request):
    """Encabezado y filas del export de fichas (compartido por CSV y Excel).

    Cubre TODAS las preguntas y subpreguntas del instrumento: identificación,
    F1 (incluidos los 'otro, ¿cuál?' y el diccionario), F2 (incluidos vínculo
    entre unidades y variables/llaves de cada fuente requerida) y F3 (bloque
    OCHA completo con sus '¿por qué?' y niveles de clasificación aplica_*).
    """
    header = [
        'UUID', 'Estado', 'Ficha inicial', 'Último paso', 'Creado', 'Actualizado',
        'Organismo', 'Suborganismo', 'Equipo', 'Responsable', 'Correo', 'Teléfono',
        # F1 · información general
        'F1 ID', 'F1 Fuente', 'F1 Descripción',
        'F1 Relaciones', 'F1 Relación (otra)',
        'F1 Instrumentos', 'F1 Instrumento (otro)',
        'F1 Mecanismos', 'F1 Mecanismo (otro)',
        'F1 Fecha inicio', 'F1 Última actualización',
        'F1 Periodicidades', 'F1 Periodicidad (otra)',
        'F1 Desagregaciones', 'F1 Desagregación (otra)',
        'F1 Registros', 'F1 Registros (fecha)',
        'F1 Tiene diccionario', 'F1 Diccionario (ubicación)', 'F1 Diccionario (URL descarga)',
        'F1 Rol', 'F1 Fuentes de origen',
        # F1 · propósito y estructura
        'F1 Decisiones que apoya',
        'F1 Unidades de análisis', 'F1 Unidad análisis (otra)',
        'F1 Variables (nombre:tipo:descripción:identificador)',
        # F1 · calidad y cobertura
        'F1 Georreferenciación', 'F1 Niveles georreferenciación', 'F1 Nivel georref (otro)',
        'F1 Problemas de calidad', 'F1 Problema calidad (otro)',
        'F1 Validación', 'F1 Validación (descripción)',
        # F2 · reto de análisis
        'F2 Reto', 'F2 Análisis propuesto', 'F2 Resultado esperado', 'F2 Decisión que apoya',
        'F2 Unidad principal', 'F2 Unidad principal (otra)',
        'F2 Unidades a relacionar', 'F2 Unidad relacionar (otra)', 'F2 Vínculo entre unidades',
        'F2 Desagregaciones', 'F2 Desagregación (otra)',
        'F2 Fuente requerida 1', 'F2 Req1 organismo custodio',
        'F2 Req1 variables necesarias (1.8)', 'F2 Req1 llave o relación (1.9)',
        'F2 Fuente requerida 2', 'F2 Req2 organismo custodio',
        'F2 Req2 variables necesarias (1.8)', 'F2 Req2 llave o relación (1.9)',
        'F2 Fuente requerida 3', 'F2 Req3 organismo custodio',
        'F2 Req3 variables necesarias (1.8)', 'F2 Req3 llave o relación (1.9)',
        'F2 Tratamientos', 'F2 Tratamiento (otro)',
        'F2 Resultado MVP', 'F2 Resultado MVP (otro)',
        'F2 Capacidades SIPDI', 'F2 Capacidad SIPDI (otra)',
        'F2 Variables a enriquecer', 'F2 Variables a enriquecer (otros)',
        'F2 Quién usará el resultado',
        'F2 Enfoques diferenciales', 'F2 Enfoque diferencial (otro)',
        # F3 · datos personales y acceso
        'F3 Datos personales',
        'F3 Variables identificadoras', 'F3 Variable identificadora (otra)',
        'F3 Datos sensibles', 'F3 Dato sensible (otro)',
        'F3 Base legítima', 'F3 Base legítima (por qué)',
        # F3 · uso y protección (OCHA)
        'F3 Proveedor base legítima', 'F3 Proveedor base legítima (por qué)',
        'F3 Finalidad específica', 'F3 Finalidad específica (por qué)',
        'F3 Datos necesarios', 'F3 Datos necesarios (por qué)',
        'F3 Conservación limitada', 'F3 Conservación limitada (por qué)',
        'F3 Salvaguardas', 'F3 Salvaguardas (por qué)',
        'F3 Transparencia', 'F3 Transparencia (por qué)',
        # F3 · divulgación y protección
        'F3 Riesgos publicación', 'F3 Riesgos publicación (por qué)',
        'F3 Terceros estándares', 'F3 Terceros estándares (por qué)',
        'F3 Riesgo de daño', 'F3 Riesgo de daño (por qué)',
        'F3 Medidas de protección', 'F3 Medida protección (otra)',
        # F3 · valoración del activo
        'F3 Susceptible a fraude', 'F3 Custodio',
        'F3 Confidencialidad', 'F3 Integridad', 'F3 Disponibilidad', 'F3 Criticidad',
        'F3 Aplica pública', 'F3 Aplica semiprivada', 'F3 Aplica privada',
        'F3 Aplica sensible', 'F3 Aplica reservada',
    ]

    def m2m(obj, campo):
        return '; '.join(o.etiqueta for o in getattr(obj, campo).all()) if obj else ''

    def disp(obj, campo):
        return getattr(obj, f'get_{campo}_display')() if obj and getattr(obj, campo) else ''

    def txt(obj, campo):
        return getattr(obj, campo) if obj else ''

    filas = []
    inscripciones = (Inscripcion.objects
                     .select_related('organismo', 'suborganismo', 'reto',
                                     'reto__unidad_principal')
                     .prefetch_related(
                         'fuentes', 'fuentes__variables', 'fuentes__fuentes_origen',
                         'fuentes__relaciones', 'fuentes__instrumentos',
                         'fuentes__mecanismos', 'fuentes__periodicidades',
                         'fuentes__desagregaciones', 'fuentes__unidades_analisis',
                         'fuentes__niveles_georref', 'fuentes__problemas_calidad',
                         'fuentes__seguridad',
                         'fuentes__seguridad__variables_identificadoras',
                         'fuentes__seguridad__datos_sensibles',
                         'fuentes__seguridad__medidas_proteccion',
                         'reto__fuentes_requeridas', 'reto__unidades_relacionar',
                         'reto__desagregaciones', 'reto__tratamientos',
                         'reto__resultados_mvp', 'reto__capacidades_sipdi',
                         'reto__enfoques_diferenciales'))
    for i in inscripciones:
        fuentes_i = list(i.fuentes.all())
        r = getattr(i, 'reto', None)
        # Una fila por fuente para conservar la asociación fuente↔atributos;
        # inscripciones sin fuentes (ficha F2) exportan una fila con F1 vacío
        for f in fuentes_i or [None]:
            s = getattr(f, 'seguridad', None) if f else None
            f_orig = '; '.join(
                f'{o.nombre} ({o.organismo_productor})' for o in f.fuentes_origen.all()) if f else ''
            f_vars = '; '.join(
                f'{v.nombre}:{v.get_tipo_display()}:{v.descripcion or ""}'
                f'{" [ID]" if v.es_identificador else ""}'
                for v in f.variables.all()) if f else ''
            f_regs = ''
            if f:
                f_regs = (str(f.num_registros) if f.num_registros is not None
                          else ('No se conoce' if f.num_registros_desconocido else ''))
            # Fuentes requeridas (1.7-1.9): hasta 3, en columnas fijas
            reqs = list(r.fuentes_requeridas.all())[:3] if r else []
            f_req = [campo for q in reqs
                     for campo in (q.nombre, q.organismo, q.variables, q.llave_relacion)]
            f_req += [''] * (12 - len(f_req))
            filas.append([
                str(i.uuid), i.get_estado_display(), i.get_ficha_inicial_display(),
                i.paso_actual,
                i.creado.strftime('%Y-%m-%d %H:%M'), i.actualizado.strftime('%Y-%m-%d %H:%M'),
                i.organismo.nombre if i.organismo else '',
                i.suborganismo.nombre if i.suborganismo else '', i.equipo_responsable,
                i.nombre_completo, i.correo, i.telefono,
                # F1 · información general
                f.pk if f else '',
                txt(f, 'nombre'), txt(f, 'descripcion'),
                m2m(f, 'relaciones'), txt(f, 'relacion_otra'),
                m2m(f, 'instrumentos'), txt(f, 'instrumento_otro'),
                m2m(f, 'mecanismos'), txt(f, 'mecanismo_otro'),
                str(f.fecha_inicio) if f and f.fecha_inicio else '',
                str(f.fecha_actualizacion) if f and f.fecha_actualizacion else '',
                m2m(f, 'periodicidades'), txt(f, 'periodicidad_otra'),
                m2m(f, 'desagregaciones'), txt(f, 'desagregacion_otra'),
                f_regs,
                str(f.num_registros_fecha) if f and f.num_registros_fecha else '',
                disp(f, 'tiene_diccionario'), txt(f, 'diccionario_ubicacion'),
                (request.build_absolute_uri(f.archivo_diccionario.url)
                 if f and f.archivo_diccionario else ''),
                disp(f, 'rol'), f_orig,
                # F1 · propósito y estructura
                txt(f, 'decisiones_apoya'),
                m2m(f, 'unidades_analisis'), txt(f, 'unidad_analisis_otra'),
                f_vars,
                # F1 · calidad y cobertura
                disp(f, 'georreferenciacion'),
                m2m(f, 'niveles_georref'), txt(f, 'nivel_georref_otro'),
                m2m(f, 'problemas_calidad'), txt(f, 'problema_calidad_otro'),
                disp(f, 'tiene_validacion'), txt(f, 'validacion_descripcion'),
                # F2 · reto de análisis
                txt(r, 'nombre'), txt(r, 'analisis_propuesto'),
                txt(r, 'resultado_esperado'), txt(r, 'decision_apoya'),
                r.unidad_principal.etiqueta if r and r.unidad_principal else '',
                txt(r, 'unidad_principal_otra'),
                m2m(r, 'unidades_relacionar'), txt(r, 'unidades_relacionar_otra'),
                txt(r, 'vinculo_unidades'),
                m2m(r, 'desagregaciones'), txt(r, 'desagregacion_otra'),
                *f_req,
                m2m(r, 'tratamientos'), txt(r, 'tratamiento_otro'),
                m2m(r, 'resultados_mvp'), txt(r, 'resultado_mvp_otro'),
                m2m(r, 'capacidades_sipdi'), txt(r, 'capacidad_sipdi_otra'),
                txt(r, 'variables_enriquecer'), txt(r, 'variables_enriquecer_otro'),
                txt(r, 'usuario_resultado'),
                m2m(r, 'enfoques_diferenciales'), txt(r, 'enfoque_diferencial_otro'),
                # F3 · datos personales y acceso
                disp(s, 'contiene_datos_personales'),
                m2m(s, 'variables_identificadoras'), txt(s, 'variable_identificadora_otra'),
                m2m(s, 'datos_sensibles'), txt(s, 'dato_sensible_otro'),
                disp(s, 'base_legitima'), txt(s, 'base_legitima_porque'),
                # F3 · uso y protección (OCHA)
                disp(s, 'proveedor_base_legitima'), txt(s, 'proveedor_base_legitima_porque'),
                disp(s, 'finalidad_especifica'), txt(s, 'finalidad_especifica_porque'),
                disp(s, 'datos_necesarios'), txt(s, 'datos_necesarios_porque'),
                disp(s, 'conservacion_limitada'), txt(s, 'conservacion_limitada_porque'),
                disp(s, 'salvaguardas'), txt(s, 'salvaguardas_porque'),
                disp(s, 'transparencia'), txt(s, 'transparencia_porque'),
                # F3 · divulgación y protección
                disp(s, 'riesgos_publicacion'), txt(s, 'riesgos_publicacion_porque'),
                disp(s, 'terceros_estandares'), txt(s, 'terceros_estandares_porque'),
                disp(s, 'riesgo_dano'), txt(s, 'riesgo_dano_porque'),
                m2m(s, 'medidas_proteccion'), txt(s, 'medida_proteccion_otra'),
                # F3 · valoración del activo
                disp(s, 'susceptible_fraude'), txt(s, 'organismo_custodio'),
                disp(s, 'confidencialidad'), disp(s, 'integridad'),
                disp(s, 'disponibilidad'), disp(s, 'criticidad'),
                disp(s, 'aplica_publica'), disp(s, 'aplica_semiprivada'),
                disp(s, 'aplica_privada'), disp(s, 'aplica_sensible'),
                disp(s, 'aplica_reservada'),
            ])
    return header, filas


def _export_variables_filas():
    """Tabla normalizada de variables F1: una fila por variable
    de cada fuente, con el contexto de inscripción y fuente.
    Compartida por el CSV dedicado y la hoja 'Variables F1' del
    Excel."""
    header = [
        'UUID inscripción', 'Organismo', 'Suborganismo', 'Responsable',
        'F1 ID', 'F1 Fuente',
        'Variable', 'Tipo de dato', 'Descripción', 'Es identificador único',
    ]
    variables = (VariableFuente.objects
                 .select_related('fuente__inscripcion__organismo',
                                 'fuente__inscripcion__suborganismo')
                 .order_by('fuente__inscripcion__organismo__nombre',
                           'fuente__nombre', 'pk'))
    filas = [[
        str(v.fuente.inscripcion.uuid),
        v.fuente.inscripcion.organismo.nombre if v.fuente.inscripcion.organismo else '',
        v.fuente.inscripcion.suborganismo.nombre if v.fuente.inscripcion.suborganismo else '',
        v.fuente.inscripcion.nombre_completo,
        v.fuente_id,
        v.fuente.nombre,
        v.nombre,
        v.get_tipo_display(),
        v.descripcion,
        'Sí' if v.es_identificador else 'No',
    ] for v in variables]
    return header, filas


@taller_staff_required
def dashboard_export(request):
    import csv
    from django.http import HttpResponse

    header, filas = _export_filas(request)
    respuesta = HttpResponse(content_type='text/csv; charset=utf-8')
    respuesta['Content-Disposition'] = 'attachment; filename="fichas_taller.csv"'
    respuesta.write('\ufeff')  # BOM para Excel
    writer = csv.writer(respuesta)
    writer.writerow(header)
    writer.writerows(filas)
    return respuesta


@taller_staff_required
def dashboard_export_excel(request):
    """Mismo export de fichas en formato Excel (.xlsx)."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header, filas = _export_filas(request)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Fichas'
    ws.append(header)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF', size=10)
        celda.fill = PatternFill('solid', fgColor='004B87')
    for fila in filas:
        ws.append(fila)
    # Ancho de columnas según contenido (tope 50)
    for idx, col in enumerate(ws.columns, 1):
        ancho = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(ancho + 2, 50)
    ws.freeze_panes = 'A2'

    # Hoja 2: variables F1 normalizadas (una fila por variable)
    header_v, filas_v = _export_variables_filas()
    ws_v = wb.create_sheet('Variables F1')
    ws_v.append(header_v)
    for celda in ws_v[1]:
        celda.font = Font(bold=True, color='FFFFFF', size=10)
        celda.fill = PatternFill('solid', fgColor='004B87')
    for fila in filas_v:
        ws_v.append(fila)
    for idx, col in enumerate(ws_v.columns, 1):
        ancho = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws_v.column_dimensions[get_column_letter(idx)].width = min(ancho + 2, 50)
    ws_v.freeze_panes = 'A2'

    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = 'attachment; filename="fichas_taller.xlsx"'
    wb.save(respuesta)
    return respuesta


@taller_staff_required
def dashboard_export_variables(request):
    """CSV dedicado de variables F1: una fila por variable."""
    import csv
    from django.http import HttpResponse

    header, filas = _export_variables_filas()
    respuesta = HttpResponse(content_type='text/csv; charset=utf-8')
    respuesta['Content-Disposition'] = 'attachment; filename="variables_fuentes_f1.csv"'
    respuesta.write('\ufeff')  # BOM para Excel
    writer = csv.writer(respuesta)
    writer.writerow(header)
    writer.writerows(filas)
    return respuesta


@taller_staff_required
def dashboard_informe(request):
    """Genera el informe de usuarios como archivo Excel descargable."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from . import wizard_taller as wz

    inscripciones = (Inscripcion.objects
                     .select_related('organismo', 'suborganismo', 'user')
                     .order_by('pk'))

    def ultimo_paso(i):
        paso = i.paso_actual
        if not paso:
            return '—'
        p = wz.PASOS_POR_SLUG.get(paso)
        return p.titulo if p else paso

    def falta(i):
        if i.estado == Inscripcion.Estado.COMPLETADA:
            return 'Completo'
        paso = i.paso_actual
        if not paso or paso == 'identificacion':
            return 'F1, F2, F3' if i.ficha_inicial == 'f1' else 'F2, F3'
        p = wz.PASOS_POR_SLUG.get(paso)
        if not p:
            return '—'
        if p.fase == 'f1':
            return 'F2, F3'
        if p.fase == 'f2':
            return 'F3'
        if p.fase == 'f3':
            return 'Confirmar'
        return '—'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Usuarios'

    # Estilos
    titulo_font = Font(bold=True, size=14, color='004b87')
    header_font = Font(bold=True, size=10, color='ffffff')
    header_fill = PatternFill(start_color='004b87', end_color='004b87', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='cccccc'),
        right=Side(style='thin', color='cccccc'),
        top=Side(style='thin', color='cccccc'),
        bottom=Side(style='thin', color='cccccc'),
    )
    completo_fill = PatternFill(start_color='c6efce', end_color='c6efce', fill_type='solid')
    pendiente_fill = PatternFill(start_color='ffeb9c', end_color='ffeb9c', fill_type='solid')
    resumen_font = Font(bold=True, size=11)

    # Título
    ws.merge_cells('A1:K1')
    ws['A1'] = 'INFORME DE USUARIOS — Taller Gestión de Datos Institucionales'
    ws['A1'].font = titulo_font
    ws['A2'] = 'Contraseña genérica para todos: TallerCali2026*'
    ws['A2'].font = Font(italic=True, color='666666')

    # Encabezados
    headers = ['pk', 'Nombre', 'Correo', 'Organismo', 'Suborganismo',
               'Equipo', 'Usuario', 'Contraseña', 'Ficha', 'Último paso', 'Falta']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Datos
    for row, i in enumerate(inscripciones, 5):
        falta_val = falta(i)
        valores = [
            i.pk,
            i.nombre_completo,
            i.correo,
            i.organismo.nombre if i.organismo else '',
            i.suborganismo.nombre if i.suborganismo else '',
            i.equipo_responsable,
            i.user.username if i.user else '',
            'TallerCali2026*',
            i.ficha_inicial.upper(),
            ultimo_paso(i),
            falta_val,
        ]
        for col, v in enumerate(valores, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = thin_border
            if col == 11:  # Columna Falta
                if falta_val == 'Completo':
                    cell.fill = completo_fill
                elif falta_val != '—':
                    cell.fill = pendiente_fill

    # Resumen
    total = inscripciones.count()
    completadas = inscripciones.filter(estado=Inscripcion.Estado.COMPLETADA).count()
    resumen_row = len(list(inscripciones)) + 6
    ws.cell(row=resumen_row, column=1, value='RESUMEN').font = resumen_font
    resumen = [
        ('Total inscripciones', total),
        ('Completadas', completadas),
        ('En progreso', total - completadas),
        ('Iniciaron en F1', inscripciones.filter(ficha_inicial='f1').count()),
        ('Iniciaron en F2', inscripciones.filter(ficha_inicial='f2').count()),
        ('Con F1', inscripciones.filter(fuentes__isnull=False).distinct().count()),
        ('Con F2', inscripciones.filter(reto__isnull=False).distinct().count()),
        ('Con F3', inscripciones.filter(fuentes__seguridad__isnull=False).distinct().count()),
    ]
    for idx, (label, val) in enumerate(resumen, 1):
        ws.cell(row=resumen_row + idx, column=1, value=label)
        ws.cell(row=resumen_row + idx, column=2, value=val).font = Font(bold=True)

    # Ajustar anchos de columna
    anchos = [6, 32, 38, 42, 28, 22, 38, 16, 8, 28, 15]
    for col, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(col)].width = ancho

    # Freeze panes
    ws.freeze_panes = 'A5'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="informe_usuarios.xlsx"'
    wb.save(response)
    return response


@xframe_options_sameorigin
def ecosistema(request):
    return render(request, 'taller/ecosistema.html')


def ecosistema_metodologia(request):
    return render(request, 'taller/ecosistema_metodologia.html')


@taller_staff_required
@xframe_options_sameorigin
def modelo_datos(request):
    """Diagrama ER interactivo: serializa el esquema de taller (+ auth.User) a JSON."""
    import json
    from django.apps import apps
    from django.contrib.auth import get_user_model
    from django.db import models as dm

    from .models_taller import OpcionTaller, FuenteRequerida
    from .models_catalogo import BaseCatalogo, CampoCatalogo, FuenteOrigen
    modelos = [
        Organismo, Suborganismo,
        OpcionTaller, Inscripcion, FuenteDatos, FuenteOrigen,
        VariableFuente, RetoAnalisis, FuenteRequerida,
        EvaluacionSeguridad, BaseCatalogo, CampoCatalogo,
        get_user_model(),
    ]
    nodos, aristas = [], []
    for m in modelos:
        meta = m._meta
        campos = []
        for f in meta.get_fields():
            if f.auto_created and not f.concrete:
                continue  # reversos
            if isinstance(f, dm.ManyToManyField):
                campos.append({'n': f.name, 't': 'M2M', 'rel': f.related_model._meta.label})
                aristas.append({'de': meta.label, 'a': f.related_model._meta.label,
                                'tipo': 'm2m', 'campo': f.name,
                                'inv': getattr(f.remote_field, 'related_name', '') or ''})
            elif f.is_relation and f.concrete:
                tipo = 'o2o' if isinstance(f, dm.OneToOneField) else 'fk'
                marca = '1:1' if tipo == 'o2o' else 'N:1'
                campos.append({'n': f.name, 't': marca, 'rel': f.related_model._meta.label})
                aristas.append({'de': meta.label, 'a': f.related_model._meta.label,
                                'tipo': tipo, 'campo': f.name,
                                'inv': getattr(f.remote_field, 'related_name', '') or ''})
            else:
                marca = 'PK' if f.primary_key else f.get_internal_type()
                campos.append({'n': f.name, 't': marca, 'rel': ''})
        nodos.append({'id': meta.label, 'nombre': meta.object_name,
                      'tabla': meta.db_table, 'doc': (m.__doc__ or '').strip(),
                      'campos': campos})
    return render(request, 'taller/modelo_datos.html', {
        'modelo_json': json.dumps({'nodos': nodos, 'aristas': aristas}),
    })
